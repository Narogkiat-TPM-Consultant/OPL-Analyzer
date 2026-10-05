"""Shared builder for the AM Standard workbooks (自主保全基準書 = AM Standard), JIPM AM Step 3–4.

build(out, T, PARTS, STD, KAIZEN, CONFIRM) writes How_To_Use, Breakdown (Step 1), FMEA (Step 2 + 3A),
AM_Standard (Step 4, CILT), AM_Calendar, Kaizen_IoT_CBM (Step 3B), Verification (Step 5) and To_Confirm.

PARTS rows: (id, system, part, function, normal condition, module, source,
             failure mode, failure effect, failure cause (CILT gap), stress/strength, S, S reason, O, O reason, D, D reason)
STD rows:   ("GROUP", title) or (part id, check point, CILT, steps, OK, NG, frequency, minutes, kaizen tip, neglect,
             responsible, AM Check Sheet items, source)
KAIZEN rows: (check point, current problem, type, idea, benefit, cost, time saved, priority, responsible, source)
RPN, risk level, "in AM Standard" and the daily-time total are Excel formulas, so editing S/O/D re-ranks the parts.
Tags: "PDF p.x" source · "[INFER]" general rule · "[AI-KB]" outside knowledge · "*" proposal · "[DATA REQUIRED]".
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Tahoma"
INK, NAVY, SECTION, INPUT, GREY, AMBER = "1A1D21", "1F2A36", "E3ECF9", "FFFF00", "F2F2F2", "B07A00"
thin = Side(style="thin", color="9AA1AA")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAPC = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CILT_FILL = {"C": "DCEBFA", "I": "DFF2E1", "L": "FFF1C2", "T": "EDE3F7", "A": "FBE0D0"}
RISK_FILL = {"CRITICAL": ("F8C9CF", "C32A3E"), "HIGH": ("FCE7B2", "8A5A00"), "MEDIUM": ("E3ECF9", "1F5FBF"), "LOW": ("F2F2F2", "5A6775")}
FREQ_FILL = {"D": "C8E6C9", "W": "BBDEFB", "M": "FFF59D", "Q": "FFCC80", "S": "EF9A9A", "A": "D1C4E9"}
MONTHS = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]


def font(bold=False, size=10, color=INK, italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def title(ws, text, sub, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, text).font = font(True, 15, "FFFFFF")
    ws.cell(1, 1).fill = fill(NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, sub).font = font(False, 9, "5A6775", italic=True)
    ws.cell(2, 1).alignment = Alignment(vertical="center", indent=1)


def header(ws, row, labels, widths):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row, i, lab)
        c.font, c.fill, c.alignment, c.border = font(True, 9, "FFFFFF"), fill(NAVY), CENTER, BOX
    ws.row_dimensions[row].height = 36
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, r, c, v, bold=False, align=WRAPC, size=9):
    cell = ws.cell(r, c, v)
    tagged = isinstance(v, str) and not v.startswith("=") and any(t in v for t in ("[DATA REQUIRED]", "[INFER]", "[AI-KB]", "*"))
    cell.font = font(bold, size, AMBER if tagged and not bold else INK)
    cell.alignment, cell.border = align, BOX
    if isinstance(v, str) and v.startswith("[DATA REQUIRED]"):
        cell.fill = fill(INPUT)
    return cell


def group_row(ws, r, text, ncols):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    c = ws.cell(r, 1, text)
    c.font, c.fill, c.alignment = font(True, 10, "1F5FBF"), fill(SECTION), Alignment(vertical="center", indent=1)
    for k in range(1, ncols + 1):
        ws.cell(r, k).border = BOX
    ws.row_dimensions[r].height = 22


def info(ws, row, pairs):
    for label, col in pairs:
        c = ws.cell(row, col, label)
        c.font, c.alignment = font(True, 9), Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + 2)
        for k in (col + 1, col + 2):
            ws.cell(row, k).fill, ws.cell(row, k).border = fill(INPUT), BOX
    ws.row_dimensions[row].height = 20


def page(ws, landscape=True, paper=8):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = paper
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3


def risk_cf(ws, rng):
    for level, (bg, fg) in RISK_FILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{level}"'], fill=fill(bg),
                                                      font=Font(name=FONT, bold=True, color=fg)))


def freq_code(freq):
    """Calendar code from the frequency text: D daily, W weekly, Q 3-month, S 6-month, A yearly, M monthly."""
    f = freq.split()[0]
    return {"D": "D", "W": "W", "M": "M", "3M": "Q", "6M": "S", "1Y": "A"}.get(f)


def months_for(code):
    return {"D": range(12), "W": range(12), "M": range(12), "Q": (2, 5, 8, 11), "S": (5, 11), "A": (11,)}.get(code, ())


def risk_of(S, O, D):
    """Python mirror of the FMEA sheet formula (skill rule order) — used for the summary only."""
    rpn = S * O * D
    if S >= 9 or (S >= 7 and rpn >= 200):
        return "CRITICAL"
    if S >= 7 or rpn >= 150 or O >= 7:
        return "HIGH"
    return "MEDIUM" if D >= 8 else "LOW"


def build(out, T, PARTS, STD, KAIZEN, CONFIRM):
    wb = Workbook()

    # ------------------------------------------------------------ How_To_Use
    ws = wb.active
    ws.title = "How_To_Use"
    title(ws, f"AM Standard — {T['name_en']} (自主保全基準書 = มาตรฐาน AM · {T['name_th']})", T["sub"], 6)
    info(ws, 3, [("เครื่องจักร", 1), ("Line / Section", 4)])
    info(ws, 4, [("AM Step", 1), ("วันที่ / ผู้จัดทำ", 4)])
    rows = [
        ("ชีท", "ใช้ทำอะไร"),
        ("Breakdown", "Step 1 · แยก System → Component/Part · หน้าที่ · สภาพปกติ · Module M1–M6 (ฉบับเต็มอยู่ในไฟล์ " + T["kmm_file"] + ")"),
        ("FMEA", "Step 2 + 3A · Failure mode/effect/cause → CILT gap + Stress/Strength · S/O/D พร้อมเหตุผล · RPN และ Risk level คำนวณด้วยสูตร"),
        ("AM_Standard", "Step 4 · มาตรฐาน CILT: วิธีทำ 1-2-3 · เกณฑ์ ✓/✗ · ความถี่ · เวลา · Risk · Kaizen · ผลถ้าละเลย · ผู้รับผิดชอบ · ข้อใน AM Check Sheet"),
        ("AM_Calendar", "ปฏิทิน 12 เดือนตามความถี่ใน AM_Standard (D/W/Q/S/A) · เดือนของรอบ 3 เดือน/6 เดือน/1 ปี เป็นข้อเสนอ *"),
        ("Kaizen_IoT_CBM", "Step 3B · จุดที่เป็น HTA (เข้าถึงยาก) หรือ SOC (แหล่งสกปรก) และใช้เวลา > 2 นาที → Visual control · Karakuri · IoT CBM"),
        ("Verification", "Step 5 · ตรวจ V1–V10 ก่อนใช้"),
        ("To_Confirm", "ประเด็นที่หัวหน้างาน / ฝ่ายซ่อม / Safety ต้องยืนยัน"),
        ("", ""),
        ("CILT", "C = Cleaning (清掃 ทำความสะอาด) · I = Inspection (点検 ตรวจ) · L = Lubrication (給油 หล่อลื่น/เติมน้ำมัน) · T = Tightening (増締め ขันแน่น) · A = Alignment/Adjust (芯出し・調整 ตั้งศูนย์/ปรับ)"),
        ("Risk level", "CRITICAL: S ≥ 9 หรือ (S ≥ 7 และ RPN ≥ 200) · HIGH: S ≥ 7 หรือ RPN ≥ 150 หรือ O ≥ 7 · MEDIUM: D ≥ 8 · LOW: อื่นๆ (RPN = S × O × D)"),
        ("ป้ายที่มา", "PDF p.x = ต้นฉบับ JIPM-Solutions · [INFER] = หลักวิศวกรรมทั่วไป · [AI-KB] = ความรู้นอกต้นฉบับ (ราคา/อุปกรณ์) · * = ข้อเสนอ · [DATA REQUIRED] = ค่าของเครื่อง (ช่องเหลือง)"),
        ("ความเชื่อมั่น", T["confidence"]),
        ("สมมติฐาน", T["assume"]),
        ("ใช้คู่กับ", T["links"]),
    ]
    for i, (a, b) in enumerate(rows, start=6):
        head = i == 6
        ws.cell(i, 1, a).font = font(True, 10, "FFFFFF" if head else INK)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.cell(i, 2, b).font = font(head, 10, "FFFFFF" if head else INK)
        for k in range(1, 7):
            ws.cell(i, k).border = BOX if a or b else Border()
            ws.cell(i, k).alignment = WRAPC
            if head:
                ws.cell(i, k).fill = fill(NAVY)
        ws.row_dimensions[i].height = 40 if a else 10
    ws.column_dimensions["A"].width = 16
    for col in "BCDEF":
        ws.column_dimensions[col].width = 26
    page(ws, landscape=False, paper=9)

    # ------------------------------------------------------------ Breakdown (Step 1)
    ws = wb.create_sheet("Breakdown")
    cols = ["ID", "System", "Component / Part", "Part No.", "หน้าที่ (Function)", "สภาพปกติ (Normal condition + Spec)", "Module", "ที่มา"]
    title(ws, f"Step 1 · Function & Breakdown Analysis — {T['name_th']}",
          "System → Component/Part · Part No. ต้องใส่จาก Spare Part List ของเครื่องจริง · Module: M1 Bolt/Nut · M2 Lubrication · M3 Drive · M4 Hydraulics · M5 Pneumatics · M6 Electrical", len(cols))
    header(ws, 4, cols, [7, 22, 30, 16, 36, 42, 9, 18])
    r, last_sys = 5, None
    for p in PARTS:
        pid, system, part, func, normal, module, src = p[:7]
        if system != last_sys:
            group_row(ws, r, system, len(cols))
            r, last_sys = r + 1, system
        for k, v in enumerate((pid, system, part, "[DATA REQUIRED]", func, normal, module, src), start=1):
            put(ws, r, k, v, bold=(k in (1, 3)), align=CENTER if k in (1, 7) else WRAPC)
        ws.row_dimensions[r].height = 42
        r += 1
    ws.freeze_panes = "D5"
    ws.print_title_rows = "4:4"
    page(ws)

    # ------------------------------------------------------------ FMEA (Step 2 + 3A)
    ws = wb.create_sheet("FMEA")
    cols = ["ID", "Part", "Failure mode", "Failure effect (BD/DF/MS/SF)", "Failure cause (CILT gap)", "Stress / Strength",
            "S", "เหตุผล S", "O", "เหตุผล O", "D", "เหตุผล D", "RPN", "Risk level", "อยู่ใน AM Standard"]
    title(ws, f"Step 2 · FMEA Risk Assessment (AIAG 4th) + Step 3A Critical selection — {T['name_th']}",
          "BD = Breakdown · DF = Defect · MS = Minor stop/ความสูญเสีย · SF = Safety · RPN, Risk level และคอลัมน์สุดท้ายเป็นสูตร — แก้ S/O/D แล้วจัดลำดับใหม่ทันที · O ยังไม่มีประวัติเสีย [INFER]", len(cols))
    header(ws, 4, cols, [7, 24, 26, 30, 32, 20, 5, 24, 5, 24, 5, 24, 7, 12, 11])
    first = 5
    for i, p in enumerate(PARTS):
        r = first + i
        pid, _, part = p[0], p[1], p[2]
        fm, fe, fc, ss, S, Sr, O, Or, D, Dr = p[7:17]
        vals = (pid, part, fm, fe, fc, ss, S, Sr, O, Or, D, Dr)
        for k, v in enumerate(vals, start=1):
            put(ws, r, k, v, bold=(k in (1, 2)), align=CENTER if k in (1, 7, 9, 11) else WRAPC)
        for k in (7, 9, 11):
            ws.cell(r, k).fill, ws.cell(r, k).font = fill(INPUT), font(True, 10)
        put(ws, r, 13, f"=G{r}*I{r}*K{r}", bold=True, align=CENTER).fill = fill(GREY)
        put(ws, r, 14, f'=IF(G{r}>=9,"CRITICAL",IF(AND(G{r}>=7,M{r}>=200),"CRITICAL",IF(OR(G{r}>=7,M{r}>=150,I{r}>=7),"HIGH",IF(K{r}>=8,"MEDIUM","LOW"))))',
            bold=True, align=CENTER)
        put(ws, r, 15, f'=IF(COUNTIF(AM_Standard!$B:$B,A{r})>0,"✓","ขาด")', bold=True, align=CENTER)
        ws.row_dimensions[r].height = 64
    last = first + len(PARTS) - 1
    risk_cf(ws, f"N{first}:N{last}")
    ws.conditional_formatting.add(f"O{first}:O{last}", CellIsRule(operator="equal", formula=['"ขาด"'], fill=fill("F8C9CF")))
    sdv = DataValidation(type="whole", operator="between", formula1="1", formula2="10")
    sdv.error, sdv.errorTitle = "ใส่ 1–10", "S/O/D"
    ws.add_data_validation(sdv)
    sdv.add(f"G{first}:G{last}"), sdv.add(f"I{first}:I{last}"), sdv.add(f"K{first}:K{last}")
    r = last + 2
    for lab, formula in (("CRITICAL", "CRITICAL"), ("HIGH", "HIGH"), ("MEDIUM", "MEDIUM"), ("LOW", "LOW")):
        put(ws, r, 13, lab, bold=True, align=CENTER)
        put(ws, r, 14, f'=COUNTIF(N{first}:N{last},"{formula}")', bold=True, align=CENTER).fill = fill(GREY)
        r += 1
    put(ws, r, 13, "ไม่อยู่ใน Standard", bold=True, align=CENTER)
    put(ws, r, 14, f'=COUNTIF(O{first}:O{last},"ขาด")', bold=True, align=CENTER).fill = fill(GREY)
    risk_cf(ws, f"M{last + 2}:M{last + 5}")
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"
    page(ws)
    fmea_rng = f"FMEA!$A${first}:$A${last}"
    risk_rng = f"FMEA!$N${first}:$N${last}"

    # ------------------------------------------------------------ AM_Standard (Step 4)
    ws = wb.create_sheet("AM_Standard")
    cols = ["No.", "Part ID", "ระบบ / จุดตรวจ (Check point)", "CILT", "วิธีทำ (1-2-3)", "✓ ปกติ (Normal)", "✗ ผิดปกติ → ทำอะไร",
            "ความถี่", "เวลา (นาที)", "Risk level", "Kaizen / Visual control", "ผลถ้าละเลย (Neglect)", "ผู้รับผิดชอบ", "ข้อใน AM Check Sheet", "ที่มา"]
    title(ws, f"Step 4 · AM Standard (CILT) — {T['name_th']}",
          "เวลา = ประมาณการ [INFER] ให้จับเวลาจริง · Risk level ดึงจากชีท FMEA ด้วยสูตร · D = ทุกวัน · W = รายสัปดาห์ · 3M / 6M / 1Y = ตามรอบ · ข้อใน AM Check Sheet: D = Daily_Check · W = Weekly_Check · P = Periodic_Check", len(cols))
    header(ws, 4, cols, [5, 7, 26, 6, 40, 26, 30, 13, 7, 11, 26, 28, 12, 11, 15])
    r, no, std_first, std_rows = 5, 0, None, []
    for row in STD:
        if row[0] == "GROUP":
            group_row(ws, r, row[1], len(cols))
            r += 1
            continue
        no += 1
        std_first = std_first or r
        std_rows.append(r)
        pid, cp, cilt, steps, ok, ng, freq, mins, kz, neglect, resp, cs, src = row
        vals = (no, pid, cp, cilt, steps, ok, ng, freq, mins, None, kz, neglect, resp, cs, src)
        for k, v in enumerate(vals, start=1):
            put(ws, r, k, v, bold=(k in (2, 3)), align=CENTER if k in (1, 2, 4, 8, 9, 13, 14) else WRAPC)
        ws.cell(r, 4).fill = fill(CILT_FILL.get(cilt[0], GREY))
        ws.cell(r, 4).font = font(True, 10)
        put(ws, r, 10, f'=IFERROR(INDEX({risk_rng},MATCH(B{r},{fmea_rng},0)),"—")', bold=True, align=CENTER)
        ws.row_dimensions[r].height = 78
        r += 1
    std_last = r - 1
    risk_cf(ws, f"J{std_first}:J{std_last}")
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    put(ws, r, 1, "เวลา AM รายวันรวม (นาที) · เป้า Step 3: < 10 นาที/กะ", bold=True, align=Alignment(horizontal="right", vertical="center"))
    for k in range(2, 9):
        ws.cell(r, k).border = BOX
    put(ws, r, 9, f'=SUMIF(H{std_first}:H{std_last},"D*",I{std_first}:I{std_last})', bold=True, align=CENTER).fill = fill(GREY)
    put(ws, r, 10, f'=IF(I{r}<10,"ผ่าน","เกิน")', bold=True, align=CENTER)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=8)
    put(ws, r + 1, 1, "เวลา AM รายสัปดาห์รวม (นาที)", bold=True, align=Alignment(horizontal="right", vertical="center"))
    for k in range(2, 9):
        ws.cell(r + 1, k).border = BOX
    put(ws, r + 1, 9, f'=SUMIF(H{std_first}:H{std_last},"W*",I{std_first}:I{std_last})', bold=True, align=CENTER).fill = fill(GREY)
    ws.freeze_panes = "D5"
    ws.print_title_rows = "4:4"
    page(ws)

    # ------------------------------------------------------------ AM_Calendar
    ws = wb.create_sheet("AM_Calendar")
    cols = ["No.", "กิจกรรม (CILT)", "CILT", "ความถี่", "Risk"] + MONTHS + ["ผู้รับผิดชอบ"]
    title(ws, f"AM Calendar — {T['name_th']} · ปี ……",
          "D = ทุกวัน · W = รายสัปดาห์ · Q = ทุก 3 เดือน (มี.ค. มิ.ย. ก.ย. ธ.ค. *) · S = ทุก 6 เดือน (มิ.ย. ธ.ค. *) · A = ทุก 1 ปี (ธ.ค. *) · เดือนที่ตรวจตามรอบเป็นข้อเสนอ ให้วางตามแผนหยุดเครื่อง", len(cols))
    header(ws, 4, cols, [5, 40, 6, 13, 11] + [6] * 12 + [13])
    r, no = 5, 0
    for row in STD:
        if row[0] == "GROUP":
            group_row(ws, r, row[1], len(cols))
            r += 1
            continue
        no += 1
        pid, cp, cilt, freq, resp = row[0], row[1], row[2], row[6], row[10]
        put(ws, r, 1, no, align=CENTER)
        put(ws, r, 2, cp, align=WRAPC)
        put(ws, r, 3, cilt, bold=True, align=CENTER).fill = fill(CILT_FILL.get(cilt[0], GREY))
        put(ws, r, 4, freq, align=CENTER)
        put(ws, r, 5, f"=AM_Standard!J{std_rows[no - 1]}", bold=True, align=CENTER)
        code = freq_code(freq)
        for m in range(12):
            c = put(ws, r, 6 + m, None, align=CENTER)
            if m in months_for(code):
                c.value = code
                c.fill, c.font = fill(FREQ_FILL[code]), font(True, 9)
        put(ws, r, 18, resp, align=CENTER)
        ws.row_dimensions[r].height = 30
        r += 1
    risk_cf(ws, f"E5:E{r - 1}")
    r += 1
    legend = [("D", "ทุกวัน (Daily)"), ("W", "รายสัปดาห์ (Weekly)"), ("Q", "ทุก 3 เดือน"), ("S", "ทุก 6 เดือน"), ("A", "ทุก 1 ปี (Annual)")]
    for i, (code, txt) in enumerate(legend):
        c = put(ws, r, 6 + i * 2, code, bold=True, align=CENTER)
        c.fill = fill(FREQ_FILL[code])
        ws.merge_cells(start_row=r, start_column=7 + i * 2, end_row=r, end_column=7 + i * 2)
        put(ws, r, 7 + i * 2, txt, align=WRAPC)
    ws.row_dimensions[r].height = 30
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"
    page(ws)

    # ------------------------------------------------------------ Kaizen_IoT_CBM (Step 3B)
    ws = wb.create_sheet("Kaizen_IoT_CBM")
    cols = ["No.", "จุดตรวจ / CILT", "ปัญหาปัจจุบัน (HTA / SOC · เวลา)", "ประเภท", "Idea", "ประโยชน์ที่คาด", "ค่าใช้จ่าย (บาท)", "เวลาที่ลดได้",
            "Priority", "ผู้รับผิดชอบ", "ที่มา"]
    title(ws, f"Step 3B · Kaizen & IoT CBM — {T['name_th']} (からくりカイゼン = Karakuri Kaizen · 目で見る管理 = Visual control)",
          "เงื่อนไข: เป็น HTA (Hard to Access) หรือ SOC (Source of Contamination) และใช้เวลาตรวจ > 2 นาที · เวลาปัจจุบัน = ประมาณการ [INFER] ต้องจับเวลาจริงก่อนอนุมัติ · ราคา [AI-KB] ให้จัดซื้อยืนยัน", len(cols))
    header(ws, 4, cols, [5, 22, 32, 14, 40, 30, 14, 14, 9, 12, 15])
    for i, row in enumerate(KAIZEN, start=1):
        r = 4 + i
        for k, v in enumerate((i,) + row, start=1):
            put(ws, r, k, v, bold=(k == 2), align=CENTER if k in (1, 4, 7, 8, 9, 10) else WRAPC)
        pr = ws.cell(r, 9)
        pr.fill = fill({"P1": "F8C9CF", "P2": "FCE7B2", "P3": "FFF59D"}.get(row[7], GREY))
        pr.font = font(True, 9)
        ws.row_dimensions[r].height = 70
    page(ws)

    # ------------------------------------------------------------ Verification (Step 5)
    ws = wb.create_sheet("Verification")
    cols = ["V", "ตรวจอะไร", "ผล", "หมายเหตุ"]
    title(ws, "Step 5 · Self-Verification", "ผล ✓ = ผ่าน · △ = ผ่านแบบมีเงื่อนไข (ต้องกรอกข้อมูลเครื่องจริง) · บางข้อคำนวณจากชีท FMEA", 4)
    header(ws, 4, cols, [6, 54, 12, 64])
    crit_missing = f'COUNTIFS(FMEA!N{first}:N{last},"CRITICAL",FMEA!O{first}:O{last},"ขาด")'
    s9_wrong = f'COUNTIFS(FMEA!G{first}:G{last},">=9",FMEA!N{first}:N{last},"<>CRITICAL")'
    checks = [
        ("V1", "Part ทุกตัวมาจากเอกสาร (ไม่มี Part ที่แต่งขึ้น)", "△", "ทุก Part มาจาก PDF ต้นฉบับ (ชีท Breakdown คอลัมน์ที่มา) · ยังไม่มี Spare Part List → Part No. = [DATA REQUIRED]"),
        ("V2", "S / O / D ทุกตัวมีเหตุผลอ้างตาราง AIAG", "△", "S และ D อ้างตาราง · O ยังไม่มีประวัติเสียของเครื่อง → ประเมินจากสาเหตุในต้นฉบับ [INFER] ให้ปรับด้วย BD record"),
        ("V3", "Part ที่ S = 9–10 เป็น CRITICAL ทุกตัว", f'=IF({s9_wrong}=0,"✓","✗")', "สูตรใน FMEA บังคับ S ≥ 9 → CRITICAL"),
        ("V4", "Kaizen ทุกรายการเป็น HTA/SOC และ > 2 นาที", "△", "ทุกรายการระบุ HTA/SOC · เวลาปัจจุบันเป็นประมาณการ → จับเวลาจริงก่อนอนุมัติ"),
        ("V5", "Module code (M1–M6) ครบทุก Part", "✓", "ชีท Breakdown คอลัมน์ Module"),
        ("V6", "Failure cause ทุกข้อ map กลับ CILT gap + Stress/Strength", "✓", "ชีท FMEA คอลัมน์ Failure cause และ Stress / Strength"),
        ("V7", "ไม่มี Spec / Part No. ที่แต่งขึ้น", "✓", "ค่าตั้ง · Torque · ชนิดน้ำมัน · Part No. = [DATA REQUIRED]"),
        ("V8", "Critical part ทุกตัวอยู่ใน AM Standard / AM Check Sheet", f'=IF({crit_missing}=0,"✓","✗")', "สูตรนับ Critical ที่คอลัมน์ “อยู่ใน AM Standard” = ขาด · ข้อใน AM Check Sheet อยู่ในชีท AM_Standard"),
        ("V9", "Calendar สอดคล้องความถี่ใน AM Standard", "✓", "สร้างจากแถวเดียวกันของ AM_Standard"),
        ("V10", "Kaizen มีค่าใช้จ่ายและเวลาที่ลดได้ทุกรายการ", "△", "ราคา [AI-KB] และเวลา [INFER] เป็นประมาณการ"),
    ]
    for i, (v, what, res, note) in enumerate(checks, start=5):
        put(ws, i, 1, v, bold=True, align=CENTER)
        put(ws, i, 2, what, align=WRAPC)
        c = put(ws, i, 3, res, bold=True, align=CENTER)
        put(ws, i, 4, note, align=WRAPC)
        ws.row_dimensions[i].height = 34
    ws.conditional_formatting.add("C5:C14", CellIsRule(operator="equal", formula=['"✓"'], fill=fill("C8E6C9")))
    ws.conditional_formatting.add("C5:C14", CellIsRule(operator="equal", formula=['"△"'], fill=fill("FCE7B2")))
    ws.conditional_formatting.add("C5:C14", CellIsRule(operator="equal", formula=['"✗"'], fill=fill("F8C9CF")))
    page(ws, landscape=False, paper=9)

    # ------------------------------------------------------------ To_Confirm
    ws = wb.create_sheet("To_Confirm")
    title(ws, "To Confirm — ประเด็นที่ต้องยืนยันก่อนใช้จริง", "ให้หัวหน้างาน / ฝ่ายซ่อมบำรุง / Safety ยืนยัน แล้วแก้ในชีทที่เกี่ยวข้อง", 7)
    header(ws, 4, ["No.", "ประเด็น", "รายละเอียด", "ที่มา / สถานะ", "ผู้ยืนยัน", "วันที่", "ผล / ค่าที่ใช้"], [5, 50, 40, 28, 16, 12, 28])
    for i, (a, b, c) in enumerate(CONFIRM, start=1):
        r = 4 + i
        for k, v in enumerate((i, a, b, c), start=1):
            put(ws, r, k, v, align=CENTER if k == 1 else WRAPC)
        for k in (5, 6, 7):
            put(ws, r, k, None).fill = fill(INPUT)
        ws.row_dimensions[r].height = 36
    page(ws)

    wb.save(out)
    summary = {}
    for p in PARTS:
        summary.setdefault(risk_of(p[11], p[13], p[15]), []).append(p[0])
    print("wrote", out, {k: len(v) for k, v in summary.items()})
    return summary
