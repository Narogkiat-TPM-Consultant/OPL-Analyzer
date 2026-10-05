#!/usr/bin/env python3
"""Build the One-Page Kaizen Sheets (Denso format) for the hydraulic unit and the pneumatic system.

Each case in kaizen_cases.py goes through the creative-idea-kaizen skill's own builder
(build_excel_template.py → One-Page Kaizen · AI Image Prompts · Principles Trace · Self-Check); the four cases of a
system are then collected into one workbook: Summary + one One-Page sheet per Kaizen + the three shared sheets.

Usage: python3 make_kaizen_sheets.py
Env:   KAIZEN_TEMPLATE = path to build_excel_template.py (default: the synced creative-idea-kaizen skill)
"""
import copy
import glob
import json
import os
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from kaizen_cases import CASES

TEMPLATE = os.environ.get("KAIZEN_TEMPLATE") or next(iter(sorted(glob.glob(
    "/root/.claude/skills/synced/*/creative-idea-kaizen/scripts/build_excel_template.py"))), None)
if not TEMPLATE or not os.path.exists(TEMPLATE):
    sys.exit("build_excel_template.py of the creative-idea-kaizen skill not found; set KAIZEN_TEMPLATE")

SYSTEMS = {
    "hyd": ("Kaizen_Sheet_Hydraulic_Unit.xlsx", "ระบบไฮดรอลิก", "AM Pillar · ระบบไฮดรอลิก", "AM_Standard_Hydraulic_Unit.xlsx"),
    "pn": ("Kaizen_Sheet_Pneumatic_System.xlsx", "ระบบลม", "AM Pillar · ระบบลม", "AM_Standard_Pneumatic_System.xlsx"),
}
F = "Tahoma"
C_DARK, C_LIGHT, C_ACCENT, C_YELLOW, C_GREY = "44546A", "EDF0F5", "C0504D", "FFF2CC", "F2F2F2"
thin = Side(style="thin", color="B0B7C3")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
LEGEND = ("คำอธิบาย: ช่องพื้นเหลือง = ช่องกรอก/แก้ได้ | ข้อมูลจาก AM Standard (ชีท Kaizen_IoT_CBM) + เอกสาร JIPM-Solutions Equipment Skills Training | "
          "Tag: [FACT] = ต้นฉบับ (ระบุหน้า) / [INFERENCE: H/M/L] = อนุมาน / [PROPOSED] = เป้าหมาย / [DATA REQUIRED] = ข้อมูลโรงงาน / [ASSUME] = ค่าตัวอย่าง")


def case_json(c, company):
    return {
        "title": c["title"], "company": company, "team": "ทีม AM (ระบุ)", "date": "DD-MMM-YY", "summary": c["summary"],
        "before": {"workflow": c["workflow"], "problem_quant": c["problem_quant"],
                   "image_note": f"[วางภาพ Before — Prompt {c['id']} B1 ในชีท 'AI Image Prompts']"},
        "ideas": c["ideas"],
        "after": {"mechanism_summary": c["mechanism"], "points": [{"name": n, "desc": d} for n, d in c["points"]],
                  "image_note": f"[วางภาพ After — Prompt {c['id']} A1 / A2 ในชีท 'AI Image Prompts']"},
        "cost": c["cost"], "result": c["result"], "yokoten": c["yokoten"],
        "prompts": [dict(zip(("image", "style", "viewpoint", "aspect", "principle", "prompt"), p)) for p in c["prompts"]],
        "principles": [dict(zip(("cat", "code", "name", "usage", "ref", "tag"), p)) for p in c["principles"]],
        "selfcheck": c["selfcheck"],
    }


def copy_sheet(src, dst):
    for row in src.iter_rows():
        for cell in row:
            d = dst.cell(row=cell.row, column=cell.column)
            if not isinstance(cell, MergedCell):
                d.value = cell.value
            if cell.has_style:
                d.font, d.fill, d.border = copy.copy(cell.font), copy.copy(cell.fill), copy.copy(cell.border)
                d.alignment, d.number_format = copy.copy(cell.alignment), cell.number_format
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[k].height = dim.height
    dst.sheet_view.showGridLines = False


def fit_rows(ws):
    """Grow row heights so wrapped Thai text in merged cells is not clipped (Excel does not auto-fit merged cells)."""
    import math
    spans = {}
    for rng in ws.merged_cells.ranges:
        spans[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1,
                                             sum(ws.column_dimensions[get_column_letter(c)].width or 8.43 for c in range(rng.min_col, rng.max_col + 1)))
    need = {}
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or isinstance(cell, MergedCell):
                continue
            nrows, width = spans.get((cell.row, cell.column), (1, ws.column_dimensions[cell.column_letter].width or 8.43))
            size = (cell.font.sz or 10) if cell.font else 10
            per_line = max(1.0, width * 1.25 * 10 / size)
            lines = sum(math.ceil(max(1, len(part)) / per_line) for part in cell.value.split("\n"))
            h = lines * size * 1.45 / nrows
            for rr in range(cell.row, cell.row + nrows):
                need[rr] = max(need.get(rr, 0), h)
    for rr, h in need.items():
        cur = ws.row_dimensions[rr].height or 15
        if h > cur:
            ws.row_dimensions[rr].height = round(h, 1)


def head(ws, labels, widths, row=1):
    for i, (h, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row, i, h)
        c.font, c.fill = Font(name=F, size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor=C_DARK)
        c.alignment, c.border = Alignment(horizontal="center", vertical="center", wrap_text=True), BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32


def cellv(ws, r, c, v, bold=False, fill=None, center=False, color="000000", size=9):
    x = ws.cell(r, c, v)
    x.font = Font(name=F, size=size, bold=bold, color=color)
    x.alignment = Alignment(horizontal="center" if center else "left", vertical="center" if center else "top", wrap_text=True)
    x.border = BOX
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    return x


def page(ws, landscape=True, paper=9):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = paper
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_system(key, tmp):
    out, sys_th, company, std_file = SYSTEMS[key]
    cases = [c for c in CASES if c["sys"] == key]
    books = {}
    for c in cases:
        j = os.path.join(tmp, c["id"] + ".json")
        x = os.path.join(tmp, c["id"] + ".xlsx")
        json.dump(case_json(c, company), open(j, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        subprocess.run([sys.executable, TEMPLATE, j, x], check=True, capture_output=True)
        books[c["id"]] = load_workbook(x)

    wb = Workbook()
    # ---------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws.merge_cells("A1:Q1")
    t = ws.cell(1, 1, f"Kaizen Sheet — {sys_th} (改善 = Kaizen · One-Page ฟอร์ม Denso)")
    t.font, t.fill = Font(name=F, size=15, bold=True, color="FFFFFF"), PatternFill("solid", fgColor=C_DARK)
    t.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:Q2")
    s = ws.cell(2, 1, f"จาก {std_file} ชีท Kaizen_IoT_CBM · เกณฑ์: HTA/SOC และใช้เวลา > 2 นาที · เวลาก่อนทำ = ประมาณการ [INFERENCE: M] ต้องจับเวลาจริง · "
                      "ครั้ง/ปี 312 = 26 วัน × 12 เดือน [ASSUME] · ช่องเหลือง = กรอกได้")
    s.font = Font(name=F, size=9, italic=True, color="5A6775")
    ws.cell(3, 1, "ค่าแรง (บาท/ชม.)").font = Font(name=F, size=9, bold=True)
    rate = cellv(ws, 3, 2, None, fill=C_YELLOW, center=True)
    ws.cell(3, 3, "[DATA REQUIRED] ใส่แล้วคอลัมน์ ‘บาท/ปี’ คำนวณให้").font = Font(name=F, size=8, italic=True, color="B07A00")
    labels = ["No.", "Kaizen", "ชื่อผลงาน", "จุดใน AM Standard", "HTA / SOC", "ก่อน (นาที/ครั้ง)", "หลัง (นาที/ครั้ง) [PROPOSED]", "ครั้ง/ปี",
              "ลดได้ (นาที/ปี)", "บาท/ปี (ค่าแรง)", "ค่าใช้จ่ายสูงสุด (บาท)", "Priority", "ผู้รับผิดชอบ", "Self-Check (/20)", "Confidence", "สถานะ", "วันเสร็จ"]
    head(ws, labels, [5, 8, 34, 16, 10, 9, 11, 8, 10, 11, 11, 8, 12, 9, 26, 12, 11], row=5)
    st = DataValidation(type="list", formula1='"วางแผน,กำลังทำ,เสร็จ,ยกเลิก"', allow_blank=True)
    ws.add_data_validation(st)
    first = 6
    for i, c in enumerate(cases):
        r = first + i
        vals = [i + 1, c["id"], c["title"], c["std"], c["kind"], c["before_min"], c["after_min"], c["occ"]]
        for k, v in enumerate(vals, start=1):
            cellv(ws, r, k, v if v is not None else "[DATA REQUIRED]", bold=(k == 2), center=(k != 3),
                  fill=C_YELLOW if k in (6, 7, 8) else None, color="B07A00" if v is None else "000000")
        cellv(ws, r, 9, f'=IFERROR(IF(OR(F{r}="",G{r}="",H{r}=""),"",(F{r}-G{r})*H{r}),"")', bold=True, center=True, fill=C_GREY)
        cellv(ws, r, 10, f'=IF(OR($B$3="",I{r}=""),"",ROUND(I{r}/60*$B$3,0))', center=True, fill=C_GREY)
        cellv(ws, r, 11, c["cost_max"] if c["cost_max"] is not None else "[DATA REQUIRED]", center=True, fill=C_YELLOW,
              color="B07A00" if c["cost_max"] is None else "000000")
        cellv(ws, r, 12, c["priority"], bold=True, center=True,
              fill={"P1": "F8C9CF", "P2": "FCE7B2", "P3": "FFF59D"}.get(c["priority"]))
        cellv(ws, r, 13, c["resp"], center=True)
        cellv(ws, r, 14, f"='Self-Check'!{get_column_letter(2 + i)}$TOTALROW", bold=True, center=True, fill=C_GREY)
        cellv(ws, r, 15, c["confidence"])
        cellv(ws, r, 16, "วางแผน", center=True, fill=C_YELLOW)
        cellv(ws, r, 17, None, center=True, fill=C_YELLOW)
        st.add(f"P{r}")
        ws.row_dimensions[r].height = 42
    last = first + len(cases) - 1
    tr = last + 1
    ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=8)
    cellv(ws, tr, 1, "รวม", bold=True, center=True)
    for k in range(2, 9):
        ws.cell(tr, k).border = BOX
    cellv(ws, tr, 9, f"=SUM(I{first}:I{last})", bold=True, center=True, fill=C_GREY)
    cellv(ws, tr, 10, f'=IF($B$3="","",SUM(J{first}:J{last}))', bold=True, center=True, fill=C_GREY)
    cellv(ws, tr, 11, f"=SUM(K{first}:K{last})", bold=True, center=True, fill=C_GREY)
    ws.merge_cells(start_row=tr + 1, start_column=1, end_row=tr + 1, end_column=8)
    cellv(ws, tr + 1, 1, "ลดได้ (ชั่วโมง/ปี)", bold=True, center=True)
    cellv(ws, tr + 1, 9, f"=ROUND(I{tr}/60,1)", bold=True, center=True, fill=C_GREY)
    ws.merge_cells(start_row=tr + 3, start_column=1, end_row=tr + 3, end_column=17)
    n = ws.cell(tr + 3, 1, "ลำดับการอ่าน: One-Page ของแต่ละ Kaizen (ชีท " + " · ".join(c["id"] for c in cases)
                + ") → AI Image Prompts (คัดลอกไปสร้างภาพ Before/After แล้ววางในช่องภาพของ One-Page) → Principles Trace → Self-Check "
                + "· Yokoten ยังไม่มีวันที่ = คะแนน 0 จนกว่าหน้างานกรอก")
    n.font, n.alignment = Font(name=F, size=9, italic=True, color="5A6775"), Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[tr + 3].height = 32
    ws.freeze_panes = "D6"
    page(ws)

    # ---------------------------------------------------------------- One-Page per Kaizen
    for c in cases:
        dst = wb.create_sheet(c["id"])
        copy_sheet(books[c["id"]]["One-Page Kaizen"], dst)
        for row in dst.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "Demo case" in cell.value:
                    cell.value = LEGEND
        fit_rows(dst)
        dst.page_setup.orientation, dst.page_setup.paperSize = "portrait", 9
        dst.page_setup.fitToWidth, dst.page_setup.fitToHeight = 1, 1
        dst.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------------------------------------------------------- shared sheets (rows of every case)
    for name, cols, widths in (("AI Image Prompts", ["Kaizen", "ภาพ", "Style", "Viewpoint", "Aspect", "หลักการที่ฝังในภาพ", "Prompt (GPT Image 2.0 — คัดลอกไปใช้ได้ทันที)"],
                                [9, 14, 18, 14, 8, 30, 95]),
                               ("Principles Trace", ["Kaizen", "หมวด", "รหัส", "หลักการ", "การใช้ใน Case นี้", "อ้างอิง (เล่ม/หน้า)", "Source Tag"],
                                [9, 16, 14, 30, 48, 26, 14])):
        ws = wb.create_sheet(name)
        head(ws, cols, widths)
        r = 2
        for c in cases:
            src = books[c["id"]][name]
            for row in src.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                cellv(ws, r, 1, c["id"], bold=True, center=True, fill=C_LIGHT)
                for k, v in enumerate(row, start=2):
                    cellv(ws, r, k, v)
                ws.row_dimensions[r].height = 110 if name == "AI Image Prompts" else 42
                r += 1
        ws.freeze_panes = "C2"
        page(ws)

    ws = wb.create_sheet("Self-Check")
    head(ws, ["รายการ"] + [c["id"] for c in cases], [44] + [10] * len(cases))
    r, sums = 2, []
    sc0 = cases[0]["selfcheck"]
    for key_r, title in (("r1", "Round 1 — ZH Police (≥ 7/8)"), ("r2", "Round 2 — Design Quality (≥ 6/7)"), ("r3", "Round 3 — Output Quality (≥ 4/5)")):
        cellv(ws, r, 1, title, bold=True, fill=C_LIGHT)
        for k in range(len(cases)):
            cellv(ws, r, 2 + k, None, fill=C_LIGHT)
        r += 1
        first_item = r
        for idx, (item, _) in enumerate(sc0[key_r]):
            cellv(ws, r, 1, item)
            for k, c in enumerate(cases):
                cellv(ws, r, 2 + k, c["selfcheck"][key_r][idx][1], center=True, fill=C_YELLOW)
            r += 1
        cellv(ws, r, 1, "รวม", bold=True)
        for k in range(len(cases)):
            col = get_column_letter(2 + k)
            cellv(ws, r, 2 + k, f"=SUM({col}{first_item}:{col}{r - 1})", bold=True, center=True, fill=C_GREY)
        sums.append(r)
        r += 2
    total_row = r
    cellv(ws, r, 1, "TOTAL (ผ่าน ≥ 17/20)", bold=True, color=C_ACCENT, size=10)
    for k in range(len(cases)):
        col = get_column_letter(2 + k)
        cellv(ws, r, 2 + k, "=" + "+".join(f"{col}{s}" for s in sums), bold=True, center=True, color=C_ACCENT, size=10)
        cellv(ws, r + 1, 2 + k, f'=IF({col}{r}>=17,"ผ่าน","ไม่ผ่าน")', bold=True, center=True)
    cellv(ws, r + 1, 1, "ผล", bold=True)
    ws.conditional_formatting.add(f"B{r + 1}:{get_column_letter(1 + len(cases))}{r + 1}",
                                  CellIsRule(operator="equal", formula=['"ผ่าน"'], fill=PatternFill("solid", fgColor="C8E6C9")))
    cellv(ws, r + 3, 1, "หมายเหตุ: Yokoten ได้ 0 จนกว่าจะกรอก Line จริง + วันที่"
          + (" · P-K4 ยังไม่มี Baseline ปริมาณน้ำ จึงยังคำนวณเงินไม่ได้" if key == "pn" else ""), size=8)
    page(ws, landscape=False)

    # resolve the Self-Check total row reference in Summary
    sm = wb["Summary"]
    for i in range(len(cases)):
        cell = sm.cell(first + i, 14)
        cell.value = cell.value.replace("$TOTALROW", f"${total_row}")
    wb.save(out)
    print("wrote", out, "sheets", wb.sheetnames)


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as tmp:
        for key in SYSTEMS:
            build_system(key, tmp)
