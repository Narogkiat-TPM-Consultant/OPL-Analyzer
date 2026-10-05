"""Shared builder for the "Know My Machine" workbooks (AM Step 0 / Step 4 General Inspection).

build(out, T, FLOW, SYSTEM, COMPONENTS, FORCED, CONDITIONS, QUIZ, CONFIRM) writes How_To_Use, System_Map,
Component_Analysis, Forced_Deterioration, Operating_Conditions, Skill_Check, Answer_Key and To_Confirm.
T holds the system-specific texts: name_en, name_th, sub, map_desc, link, part_col, map_title, map_sub,
flow_label, videos, module. Text tags: "PDF p.x" source · "[INFER]" general rule · "*" proposal ·
"[DATA REQUIRED]" machine value (yellow).
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Tahoma"
INK, NAVY, SECTION, INPUT, GREY, AMBER = "1A1D21", "1F2A36", "E3ECF9", "FFFF00", "F2F2F2", "B07A00"
thin = Side(style="thin", color="9AA1AA")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAPC = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def font(bold=False, size=10, color=INK, italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)

# ---------------------------------------------------------------- helpers
def title(ws, text, sub, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, text).font = font(True, 15, "FFFFFF")
    ws.cell(1, 1).fill = fill(NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, sub).font = font(False, 9, "5A6775", italic=True)
    ws.cell(2, 1).alignment = Alignment(vertical="center", indent=1)


def header(ws, row, labels, widths):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row, i, lab)
        c.font, c.fill, c.alignment, c.border = font(True, 9, "FFFFFF"), fill(NAVY), CENTER, BOX
    ws.row_dimensions[row].height = 34
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, r, c, v, bold=False, align=WRAP, size=9):
    cell = ws.cell(r, c, v)
    tagged = isinstance(v, str) and ("[DATA REQUIRED]" in v or "[INFER]" in v or "*" in v)
    cell.font = font(bold, size, AMBER if tagged and not bold else INK)
    cell.alignment, cell.border = align, BOX
    if isinstance(v, str) and v.startswith("[DATA REQUIRED]"):
        cell.fill = fill(INPUT)
    return cell


def info(ws, row, pairs):
    for label, col in pairs:
        c = ws.cell(row, col, label)
        c.font, c.alignment = font(True, 9), Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + 2)
        for k in (col + 1, col + 2):
            ws.cell(row, k).fill, ws.cell(row, k).border = fill(INPUT), BOX
    ws.row_dimensions[row].height = 20


def page(ws, landscape=True, paper=8):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = paper
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3



def build(out, T, FLOW, SYSTEM, COMPONENTS, FORCED, CONDITIONS, QUIZ, CONFIRM):
    wb = Workbook()

    # ---------------------------------------------------------------- How_To_Use
    ws = wb.active
    ws.title = "How_To_Use"
    title(ws, f"Know My Machine — {T['name_en']} (設備を知る = รู้จักเครื่องจักรของเรา · {T['name_th']})",
          T["sub"], 6)
    info(ws, 3, [("เครื่องจักร", 1), ("Line / แผนก", 4)])
    info(ws, 4, [("ผู้จัดทำ", 1), ("วันที่", 4)])
    rows = [
        ("ชีท", "ใช้ทำอะไร"),
        ("System_Map", T["map_desc"]),
        ("Component_Analysis", "แต่ละชิ้นส่วน: หน้าที่ · โครงสร้าง · หลักการ · สภาพที่ถูกต้อง · การเสื่อม (ปกติ/เร่ง) · ผลถ้าปล่อยไว้ · จุดตรวจ AM · งาน PM · Safety"),
        ("Forced_Deterioration", f"6 สาเหตุของการเสีย (6 Factors of failures) → ตัวอย่างใน{T['name_th']} → มาตรการ → ข้อตรวจใน AM Check Sheet"),
        ("Operating_Conditions", "ค่ามาตรฐานและรอบการดูแล (ตามต้นฉบับ) + ช่องกรอกค่าของเครื่องนี้"),
        ("Skill_Check", "แบบทดสอบ 10 ข้อหลังดูวิดีโอ + ตารางผลรายคน (ระดับทักษะ 1–4) · เฉลยอยู่ชีท Answer_Key"),
        ("To_Confirm", "ประเด็นที่หัวหน้างาน/ฝ่ายซ่อม/Safety ต้องยืนยันก่อนใช้จริง"),
        ("", ""),
        ("ป้ายที่มา", "PDF p.x = จากต้นฉบับ · [INFER] = หลักวิศวกรรมทั่วไป ไม่อยู่ในต้นฉบับ · * = ข้อเสนอ ให้หัวหน้างานยืนยัน · [DATA REQUIRED] = ค่าของเครื่อง (ช่องเหลือง)"),
        ("ความเชื่อมั่น (Confidence)", "HIGH: ข้อที่มี PDF p.x · MEDIUM: ข้อ [INFER] · ต้องกรอก [DATA REQUIRED] ก่อนใช้กับเครื่องจริง"),
        ("วิธีใช้", "1) ทีมเดินดูเครื่องจริงพร้อมชีท System_Map กรอกตำแหน่ง · 2) หัวหน้าทีมสอนด้วยวิดีโอ + Component_Analysis · 3) ทำ Skill_Check · 4) ติด Visual control ตามจุดตรวจ"),
        ("เชื่อมกับ", T["link"]),
    ]
    for i, (a, b) in enumerate(rows, start=6):
        head = i == 6
        ws.cell(i, 1, a).font = font(True, 10, "FFFFFF" if head else INK)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.cell(i, 2, b).font = font(head, 10, "FFFFFF" if head else INK)
        for k in range(1, 7):
            ws.cell(i, k).border = BOX if a or b else Border()
            ws.cell(i, k).alignment = WRAPC
            if head:
                ws.cell(i, k).fill = fill(NAVY)
        ws.row_dimensions[i].height = 36 if a else 10
    ws.column_dimensions["A"].width = 22
    for col in "BCDEF":
        ws.column_dimensions[col].width = 24
    page(ws, landscape=False, paper=9)

    # ---------------------------------------------------------------- System_Map
    ws = wb.create_sheet("System_Map")
    cols = [T["part_col"], "กลุ่ม", "ชิ้นส่วน", "หน้าที่", "ตำแหน่ง / Tag No. บนเครื่อง", "จำนวน", "OPL · วิดีโอ", "ที่มา"]
    title(ws, T["map_title"], T["map_sub"], len(cols))
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols))
    c = ws.cell(3, 1, T["flow_label"] + FLOW)
    c.font, c.fill, c.alignment = font(True, 10, "1F5FBF"), fill(SECTION), Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[3].height = 30
    header(ws, 5, cols, [9, 26, 30, 44, 30, 9, 22, 16])
    for r, row in enumerate(SYSTEM, start=6):
        order, group, comp, func, opl, src = row
        for k, v in enumerate((order, group, comp, func), start=1):
            put(ws, r, k, v, bold=(k == 3), align=CENTER if k == 1 else WRAPC)
        for k in (5, 6):
            put(ws, r, k, None, align=CENTER).fill = fill(INPUT)
        put(ws, r, 7, opl, align=WRAPC)
        put(ws, r, 8, src, align=WRAPC)
        ws.row_dimensions[r].height = 34
    ws.freeze_panes = "C6"
    ws.print_title_rows = "5:5"
    page(ws)

    # ---------------------------------------------------------------- Component_Analysis
    ws = wb.create_sheet("Component_Analysis")
    cols = ["ชิ้นส่วน", "หน้าที่ (Function)", "โครงสร้าง / ชิ้นส่วนสำคัญ", "หลักการทำงาน", "สภาพที่ถูกต้อง (Basic condition)",
            "การเสื่อม (ปกติ / เร่ง = Forced)", "ถ้าปล่อยไว้ → ผลกระทบ", "จุดตรวจ AM (วิธี · รอบ)", "งาน PM", "Safety", "OPL · วิดีโอ", "ที่มา"]
    title(ws, f"Know My Machine — Component Analysis · {T['name_th']}",
          "อ่านจากซ้ายไปขวา: รู้หน้าที่ → รู้สภาพที่ถูกต้อง → รู้ว่าเสื่อมอย่างไร → รู้ว่าต้องตรวจอะไร · ข้อความสีส้ม = [INFER] / * / [DATA REQUIRED]", len(cols))
    header(ws, 4, cols, [17, 24, 26, 30, 26, 30, 26, 32, 24, 22, 17, 15])
    for r, row in enumerate(COMPONENTS, start=5):
        for k, v in enumerate(row, start=1):
            cell = put(ws, r, k, v, bold=(k == 1))
            if k == 1:
                cell.fill = fill(SECTION)
            if k == 10 and v != "—":
                cell.font = font(True, 9, "C32A3E" if "[INFER]" not in v and "*" not in v else AMBER)
        ws.row_dimensions[r].height = 90
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"
    page(ws)

    # ---------------------------------------------------------------- Forced_Deterioration
    ws = wb.create_sheet("Forced_Deterioration")
    cols = ["6 สาเหตุของการเสีย (6 Factors of failures)", f"ตัวอย่างใน{T['name_th']}", "ที่มา", "มาตรการ (Countermeasure)", "ข้อตรวจใน AM Check Sheet"]
    title(ws, "Forced Deterioration — การเสื่อมแบบเร่ง (強制劣化 = เสื่อมเพราะคนหรือสภาพแวดล้อม)",
          "สมการ: การเสีย = ภาระ (Stress) > ความแข็งแรง (Strength) · 6 สาเหตุ → 6 มาตรการ (JIPM AM)", len(cols))
    header(ws, 4, cols, [34, 56, 16, 40, 30])
    for r, row in enumerate(FORCED, start=5):
        for k, v in enumerate(row, start=1):
            cell = put(ws, r, k, v, bold=(k == 1), align=WRAPC)
            if k == 1:
                cell.fill = fill(SECTION)
        ws.row_dimensions[r].height = 52
    page(ws)

    # ---------------------------------------------------------------- Operating_Conditions
    ws = wb.create_sheet("Operating_Conditions")
    cols = ["หัวข้อ", "ค่ามาตรฐาน / รอบ (ต้นฉบับ)", "ที่มา", "ค่าของเครื่องนี้", "ผู้ยืนยัน"]
    title(ws, "Operating Conditions — ค่ามาตรฐานและรอบการดูแล", "ช่องสีเหลือง = กรอกค่าของเครื่องนี้ (ดู Catalog ผู้ผลิต / ฝ่ายซ่อม)", len(cols))
    header(ws, 4, cols, [34, 60, 14, 28, 16])
    for r, (a, b, src) in enumerate(CONDITIONS, start=5):
        put(ws, r, 1, a, bold=True, align=WRAPC)
        put(ws, r, 2, b, align=WRAPC)
        put(ws, r, 3, src, align=WRAPC)
        for k in (4, 5):
            put(ws, r, k, None, align=WRAPC).fill = fill(INPUT)
        ws.row_dimensions[r].height = 32
    page(ws, landscape=False, paper=9)

    # ---------------------------------------------------------------- Skill_Check
    ws = wb.create_sheet("Skill_Check")
    title(ws, f"Skill Check — แบบทดสอบความรู้{T['name_th']} ({len(QUIZ)} ข้อ)", f"ทำหลังดูวิดีโอ {T['videos']} · เฉลยอยู่ชีท Answer_Key · ผ่าน ≥ 8/10 * (ข้อเสนอ)", 6)
    info(ws, 3, [("ชื่อ", 1), ("วันที่", 4)])
    header(ws, 5, ["ข้อ", "คำถาม", "คำตอบ", "", "", "ถูก (1/0)"], [6, 52, 30, 14, 14, 10])
    for i, (q, _, ref) in enumerate(QUIZ, start=1):
        r = 5 + i
        put(ws, r, 1, i, align=CENTER)
        put(ws, r, 2, q + f"  ({ref.split(' · ')[0]})", align=WRAPC)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        for k in (3, 4, 5):
            ws.cell(r, k).border = BOX
        put(ws, r, 6, None, align=CENTER).fill = fill(INPUT)
        ws.row_dimensions[r].height = 40
    score_row = 6 + len(QUIZ)
    ws.merge_cells(start_row=score_row, start_column=1, end_row=score_row, end_column=5)
    put(ws, score_row, 1, "คะแนนรวม", bold=True, align=Alignment(horizontal="right", vertical="center"))
    for k in range(2, 6):
        ws.cell(score_row, k).border = BOX
    sc = put(ws, score_row, 6, f"=SUM(F6:F{score_row - 1})", bold=True, align=CENTER)
    sc.fill = fill(GREY)
    dv01 = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    ws.add_data_validation(dv01)
    dv01.add(f"F6:F{score_row - 1}")

    # team result table
    t0 = score_row + 2
    ws.merge_cells(start_row=t0, start_column=1, end_row=t0, end_column=6)
    ws.cell(t0, 1, f"ผลรายคน (Skill matrix · {T['module']})").font = font(True, 11, "1F5FBF")
    header(ws, t0 + 1, ["No.", "ชื่อ", "คะแนน (0–10)", "ผล", "ระดับทักษะ (1–4)", "ผู้ประเมิน"], [6, 52, 30, 14, 14, 10])
    lv = DataValidation(type="list", formula1='"1 รู้,2 ทำได้,3 มั่นใจ,4 สอนได้"', allow_blank=True)
    lv.prompt, lv.promptTitle = "1 รู้ (Know) · 2 ทำได้ (Can manage) · 3 มั่นใจ (Confident) · 4 สอนได้ (Can teach)", "ระดับทักษะ"
    ws.add_data_validation(lv)
    for n in range(1, 11):
        r = t0 + 1 + n
        put(ws, r, 1, n, align=CENTER)
        for k in (2, 3, 5, 6):
            put(ws, r, k, None, align=CENTER).fill = fill(INPUT)
        res = put(ws, r, 4, f'=IF(C{r}="","",IF(C{r}>=8,"ผ่าน","ทบทวน"))', bold=True, align=CENTER)
        res.fill = fill(GREY)
        lv.add(f"E{r}")
        ws.row_dimensions[r].height = 22
    r = t0 + 12
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, f'=COUNTIF(D{t0 + 2}:D{t0 + 11},"ผ่าน")&" คนผ่าน จาก "&COUNT(C{t0 + 2}:C{t0 + 11})&" คนที่ทดสอบ"').font = font(True, 10)
    ws.cell(r, 1).alignment = Alignment(horizontal="right", vertical="center")
    page(ws, landscape=False, paper=9)

    # ---------------------------------------------------------------- Answer_Key
    ws = wb.create_sheet("Answer_Key")
    title(ws, "Answer Key — เฉลย Skill Check", "สำหรับผู้สอน · ไม่ต้องพิมพ์แจกผู้เรียน", 4)
    header(ws, 4, ["ข้อ", "คำถาม", "เฉลย", "อ้างอิง"], [6, 50, 56, 22])
    for i, (q, a, ref) in enumerate(QUIZ, start=1):
        r = 4 + i
        put(ws, r, 1, i, align=CENTER)
        put(ws, r, 2, q, align=WRAPC)
        put(ws, r, 3, a, bold=True, align=WRAPC)
        put(ws, r, 4, ref, align=WRAPC)
        ws.row_dimensions[r].height = 34
    page(ws, landscape=False, paper=9)

    # ---------------------------------------------------------------- To_Confirm
    ws = wb.create_sheet("To_Confirm")
    title(ws, "To Confirm — ประเด็นที่ต้องยืนยันก่อนใช้จริง", "ให้หัวหน้างาน / ฝ่ายซ่อมบำรุง / Safety ยืนยัน แล้วแก้ในชีทที่เกี่ยวข้อง", 7)
    header(ws, 4, ["No.", "ประเด็น", "รายละเอียด", "ที่มา / สถานะ", "ผู้ยืนยัน", "วันที่", "ผล / ค่าที่ใช้"], [5, 50, 36, 30, 16, 12, 28])
    for i, (a, b, c) in enumerate(CONFIRM, start=1):
        r = 4 + i
        for k, v in enumerate((i, a, b, c), start=1):
            put(ws, r, k, v, align=CENTER if k == 1 else WRAPC)
        for k in (5, 6, 7):
            put(ws, r, k, None).fill = fill(INPUT)
        ws.row_dimensions[r].height = 36
    page(ws)

    wb.save(out)
    print("wrote", out)
