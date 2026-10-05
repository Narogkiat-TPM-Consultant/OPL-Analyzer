#!/usr/bin/env python3
"""Build the "Know My Machine" workbook for a pneumatic (compressed-air) system, for AM Step 0 / Step 4 Module 4-4
(Pneumatics), from the JIPM-Solutions deck Equipment Skills Training – Pneumatic (OPL 5'-A-1 … 5'-C-5, PDF p.35–49),
taught in videos EP22–EP31.

Usage: python3 make_know_my_machine_pneumatic.py [output.xlsx]

Tags: "PDF p.x" = from the deck · "[INFER]" = general engineering knowledge, not in the deck · "*" = proposal for the
supervisor to confirm · "[DATA REQUIRED]" = a value of this machine to fill in.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "Know_My_Machine_Pneumatic_System.xlsx"
FONT = "Tahoma"
INK, NAVY, SECTION, INPUT, GREY, AMBER = "1A1D21", "1F2A36", "E3ECF9", "FFFF00", "F2F2F2", "B07A00"
thin = Side(style="thin", color="9AA1AA")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAPC = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def font(bold=False, size=10, color=INK, italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


# ---------------------------------------------------------------- content
FLOW = "Compressor (แหล่งลม) → Stop valve → Air filter → Regulator + เกจ → Lubricator → Solenoid valve → Speed controller → Air cylinder · ไอเสียออกที่ Exhaust port / Silencer"

# (order, group, component, function, OPL · video, source)
SYSTEM = [
    ("①", "แหล่งลม (Air source)", "Compressor", "ผลิตลมอัด ต่อท่อไปทั้งโรงงานได้ง่าย", "5'-A-1, 5'-A-2 · EP22", "PDF p.35–36"),
    ("—", "ระบบท่อ (Piping)", "ท่อ · ข้อต่อ · Stop valve · สายยาง", "ส่งลมจาก Compressor ถึงเครื่อง", "5'-B-1, 5'-C-2 · EP26, EP29", "PDF p.41, p.45"),
    ("②", "อุปกรณ์ประกอบ (Accessories)", "Air filter (ตัวกรองลม)", "กรองฝุ่นและน้ำออกจากลมโรงงาน · ตัวแรกของชุด 3 ตัว (FRL)", "5'-A-4 · EP23", "PDF p.38, p.48"),
    ("③", "ส่วนควบคุม (Controller) · Pressure", "Regulator (ตัวปรับแรงดัน)", "ลดแรงดันให้คงที่ตามที่ใช้ แม้แรงดันขาเข้าเปลี่ยน", "5'-A-5 · EP24", "PDF p.39, p.48"),
    ("②", "อุปกรณ์ประกอบ (Accessories)", "Pressure gauge (เกจแรงดัน)", "ยืนยันค่าแรงดันที่ตั้ง", "5'-A-5 · EP24", "PDF p.36, p.39"),
    ("②", "อุปกรณ์ประกอบ (Accessories)", "Lubricator (ตัวจ่ายน้ำมัน)", "ทำน้ำมันเป็นละอองผสมลม ไปหล่อลื่นวาล์วและกระบอกสูบ", "5'-A-6 · EP25", "PDF p.40, p.48"),
    ("③", "ส่วนควบคุม (Controller) · Direction", "Solenoid valve / Direction change valve", "เปลี่ยนทิศทางลมเข้า–ออกกระบอกสูบ", "5'-A-2, 5'-C-2 · EP22, EP29", "PDF p.36, p.45"),
    ("③", "ส่วนควบคุม (Controller) · Flow", "Speed controller (วาล์วปรับความเร็ว)", "ปรับอัตราไหล → ความเร็วกระบอกสูบ", "5'-A-2, 5'-C-2 · EP22, EP29", "PDF p.36, p.45"),
    ("③", "ส่วนควบคุม (Controller) · Pressure", "Relief valve (วาล์วระบาย)", "วาล์วควบคุมแรงดัน (Safety valve)", "5'-C-3 · EP29", "PDF p.36–37, p.46"),
    ("④", "ส่วนทำงาน (Actuator)", "Air cylinder (กระบอกสูบลม)", "เปลี่ยนแรงดันลมเป็นการเคลื่อนที่", "5'-A-2, 5'-C-2 · EP22, EP29", "PDF p.36, p.45"),
    ("②", "อุปกรณ์ประกอบ (Accessories)", "Silencer / Exhaust port", "ทางลมระบายของวาล์ว · จุดดูน้ำมันที่ออกมา", "5'-A-6, 5'-C-2 · EP25, EP29", "PDF p.36, p.40, p.45"),
    ("—", "สายลม (Air tube)", "Air tube + Binder / Lock band", "ส่งลมระหว่างอุปกรณ์ · รัดกันเสียดสีและให้ทำความสะอาดง่าย", "5'-C-4 · EP30", "PDF p.47"),
]

# (component, function, key parts, how it works, correct condition, deterioration + forced cause, effect if left,
#  AM check (method · interval), PM work, safety, OPL · video, source)
COMPONENTS = [
    ("Air filter\n(ตัวกรองลม)",
     "กรองสิ่งปนเปื้อน (น้ำมัน · น้ำ · ฝุ่น · สนิม) ออกจากลมโรงงาน ก่อนถึง Solenoid valve และกระบอกสูบ",
     "Deflector · Filter element · Baffle plate (ต้นฉบับ “Bubble plate”) · ถ้วยใส (Bowl) · Drain (Manual / Auto)",
     "ลมเข้า → Deflector ทำให้หมุนเหวี่ยง สิ่งปนเปื้อนชิดผนังแล้วตกก้นถ้วย → ลมผ่าน Element ไป OUT · Baffle plate กันสิ่งที่ตกแล้วไม่ให้ปนกลับ",
     "ไม่มีน้ำขังในถ้วย · Element ไม่ตัน · ถ้วยไม่แตกร้าว",
     "น้ำและสิ่งสกปรกสะสม → Element ตัน\nเร่ง: ไม่ระบาย Drain · กลางคืนอุณหภูมิลดจึงเกิดน้ำ · สนิมในท่อทำให้ Filter ตันเร็ว",
     "น้ำ/สนิมไปถึงวาล์ว–กระบอกสูบ → ทำงานผิดปกติ ลมรั่ว · Element ตัน → เข็มเกจ Regulator แกว่ง",
     "น้ำขังในถ้วย: ตา · ทุกวัน → เปิด Drain\nเข็มเกจแกว่งขณะเดิน: ตา · ทุกวัน → เปลี่ยน Element\nAuto drain ทำงาน: 3 เดือน",
     "1 ปี: Case แตกร้าว · Element ตัน\nElement โลหะ: ล้างน้ำมันก๊าด + แปรงไนลอน · Element อัลลอย (สีเงิน): เปลี่ยนใหม่ ห้ามใช้ซ้ำ",
     "ปล่อยแรงดันค้างก่อนถอดถ้วย — ถ้วยอาจหลุดกระเด็น",
     "5'-A-4, 5'-C-5 · EP23, EP31", "PDF p.38, p.42, p.44–46, p.48–49"),
    ("Regulator\n(ตัวปรับแรงดัน)",
     "ลดแรงดันลมโรงงานให้คงที่ตามที่ใช้ แม้แรงดันขาเข้าเปลี่ยน",
     "Handle · Lock nut · Adjust spring · Diaphragm · Stem · Valve unit",
     "หมุน Handle กดสปริง → Stem ดันวาล์วเปิด ลมไหล · ถ้าแรงดันขาออกสูงกว่าค่าตั้ง Diaphragm ดึงขึ้น → วาล์วปิด ลมหยุด · สมดุลแรงสปริงกับแรงดันบน Diaphragm",
     "ตั้งค่าแล้วขัน Lock nut · เกจชี้ค่าตั้ง (แถบเขียว)\nขัน Handle → แรงดันเพิ่ม · คลาย → ลด",
     "Diaphragm เสียหาย · สปริงสนิม/ล้า/หัก · Packing บาก\nเร่ง: Lock nut หลวม → Handle หมุนเองจากแรงสั่นของเครื่อง",
     "ค่าตั้งแรงดันเคลื่อน → แรงของกระบอกสูบไม่ตรงที่ต้องการ [INFER]",
     "Lock nut หลวม: ตา · ทุกวัน → ขันให้แน่น\nลองหมุนปรับดูเกจ: 3 เดือน",
     "1 ปี: Diaphragm · สปริง · Packing",
     "—",
     "5'-A-5, 5'-C-5 · EP24, EP31", "PDF p.39, p.45–46, p.49"),
    ("Pressure gauge\n(เกจแรงดัน)",
     "ยืนยันค่าแรงดันที่ตั้ง",
     "กระจก · เข็ม · สเกล · แถบเขียวที่ค่าตั้ง",
     "—",
     "ขณะเดิน: ชี้ในช่วงใช้งานปกติ (แถบเขียว) · ปล่อยแรงดันแล้วชี้ 0",
     "กระจกแตก · เข็มผิดปกติ · ค่าเพี้ยน (ต้นฉบับ “confusion” = 狂い)",
     "อ่านค่าแรงดันผิด → ตั้งแรงดันผิด [INFER]",
     "ชี้ในช่วงปกติ: ตา · ทุกวัน\nกระจก/เข็ม: ตา → เปลี่ยนเกจ\nปล่อยแรงดันแล้วชี้ 0: 3 เดือน",
     "1 ปี: ความเที่ยง (Accuracy)",
     "—",
     "5'-A-5, 5'-C-1 · EP24, EP28", "PDF p.39, p.44–46"),
    ("Lubricator\n(ตัวจ่ายน้ำมัน)",
     "ทำน้ำมันเป็นละอองผสมลม ส่งไปหล่อลื่นอุปกรณ์ (ยกเว้นแบบ Oil-free)",
     "Oil drop window (Sight glass) · Drop adjust screw · Oil cap · Bowl\nแบบ All-amount / Selective",
     "ลมเร็วผ่านวาล์วใต้หน้าต่าง → แรงดันต่ำ → น้ำมันถูกดันขึ้นท่อ หยดในหน้าต่าง → ลมเร็วตีเป็นละออง → OUT (หลักการเดียวกับกระบอกฉีดฝอย)",
     "หยด 2–6 หยด/นาที (ค่าจริงขึ้นกับอุปกรณ์ ปรึกษาฝ่ายซ่อม) · น้ำมันใส",
     "น้ำปน → น้ำมันขุ่นขาว · น้ำมันหมด/ไม่หยด\nเร่ง: ไม่เติมตามรอบ · ปรับหยดเอง · จุดรั่ว (น้ำมันลดเร็ว)",
     "หล่อลื่นไม่พอ → เสียดทานเพิ่ม (กระบอกสูบไม่นิ่ง แรงไม่พอ วาล์วสลับไม่ดี) · ซีลสึก → ลมรั่ว · สนิมติดในวาล์ว/ลูกสูบ",
     "นับหยด 1 นาทีขณะเดิน: ทุกวัน\nน้ำมันขุ่นขาว: ทุกวัน → เปลี่ยนน้ำมันใหม่\nเติมน้ำมัน: สัปดาห์ละครั้ง (เดินหนัก ทุก 3 วัน)",
     "ตั้งค่าหยดตามอุปกรณ์และสภาพใช้งาน",
     "ปล่อยแรงดันก่อนเปิด Oil cap — แบบ Selective น้ำมันพุ่งออก",
     "5'-A-6, 5'-B-3, 5'-C-5 · EP25, EP27, EP31", "PDF p.40, p.43–44, p.49"),
    ("Solenoid valve\n(Direction change valve)",
     "เปลี่ยนทิศทางลมเข้า–ออกกระบอกสูบ",
     "Coil · Magnet core · Spool · Spring · O-ring · Piston · Diaphragm · Cam lever · Exhaust port",
     "—",
     "สลับทิศได้ราบรื่น · Exhaust port มีละอองน้ำมันพอดี · ไม่มีลมรั่ว",
     "Valve seat มีฝุ่น/รอยบาก · สปริงหัก/สนิม · Spool บาก/สึก · O-ring เสียรูป/แข็ง · ฉนวน Coil\nเร่ง: สนิม/น้ำจากท่อติดในวาล์ว · หล่อลื่นไม่พอ",
     "วาล์วสลับไม่ดี ทำงานผิดปกติ · ลมรั่วจากวาล์ว",
     "Exhaust port ดูละอองน้ำมัน: 3 เดือน (น้อย → ตรวจ Lubricator · มาก → ปรับลด · รั่ว → ดู Valve seat/สปริง)\nลมรั่วที่ Exhaust port: 3 เดือน",
     "1 ปี: Coil · Core · สปริง · Piston · Spool · Diaphragm · O-ring · Cam lever",
     "ตัดไฟ + ปล่อยลมค้างก่อนถอด *",
     "5'-C-2, 5'-C-3 · EP29", "PDF p.42–43, p.45–46"),
    ("Speed controller\n(Speed adjust valve)",
     "ปรับอัตราไหล → ความเร็วกระบอกสูบ",
     "Needle valve · Seat packing",
     "—",
     "ความเร็วการทำงานเหมือนทุกครั้ง · ปรับแล้วความเร็วเปลี่ยนตาม",
     "Needle valve บาก · Seat packing เสียหาย",
     "ความเร็วเปลี่ยนเอง / ปรับไม่ได้ [INFER]",
     "ความเร็วการทำงาน: ตา + นาฬิกา · ทุกวัน\nลองปรับความเร็ว: 3 เดือน",
     "1 ปี: Needle valve · Seat packing",
     "—",
     "5'-C-1, 5'-C-2, 5'-C-3 · EP28, EP29", "PDF p.44–46"),
    ("Air cylinder\n(กระบอกสูบลม)",
     "เปลี่ยนแรงดันลมเป็นการเคลื่อนที่ (Actuator)",
     "Tube · Piston · Rod · O-ring · Packing",
     "—",
     "เคลื่อนที่ราบรื่น ความเร็วปกติ · ไม่รั่วที่ก้านสูบ",
     "Tube บาก/เสียรูป · O-ring/Packing บาก/สึก\nเร่ง: หล่อลื่นไม่พอ · สนิมจากท่อติดในส่วนเลื่อน",
     "ไม่นิ่ง · แรงไม่พอ · ลมรั่ว",
     "ความเร็วการทำงาน: ทุกวัน\nลมรั่วที่ก้านสูบ: 3 เดือน",
     "1 ปี: ผิวใน Tube · O-ring · Packing",
     "กระบอกสูบอาจขยับเองตอนจ่ายลมกลับ — อยู่นอกระยะเคลื่อนที่ [INFER]",
     "5'-B-3, 5'-C-2, 5'-C-3 · EP27, EP29", "PDF p.42–43, p.45–46"),
    ("Relief valve\n(วาล์วระบาย)",
     "วาล์วควบคุมแรงดัน (Safety valve)",
     "Valve seat · Diaphragm · Spring",
     "—",
     "ไม่มีรอยบาก ไม่เสียหาย",
     "Valve seat บาก · Diaphragm เสียหาย · สปริงสนิม/หัก",
     "[DATA REQUIRED] ตามตำแหน่งติดตั้งในวงจรของเครื่อง",
     "ลมรั่วที่ Relief hole ของ Regulator: 3 เดือน",
     "1 ปี: Valve seat · Diaphragm · สปริง",
     "—",
     "5'-C-2, 5'-C-3 · EP29", "PDF p.36–37, p.45–46"),
    ("ระบบท่อ · ข้อต่อ\n(Piping)",
     "ส่งลมจาก Compressor ไปทั่วโรงงานถึงเครื่อง",
     "ท่อ · ข้อต่อเชื่อม · ข้อต่อเกลียว · ข้อต่อหน้าแปลน · Stop valve · สายยาง · ที่ยึดท่อ",
     "ลมเสียดสีผนังท่อ → แรงดันตกที่ปลายทาง · ท่อยาวหรือเล็กเกิน → แรงดันตกมากขึ้น",
     "ไม่รั่ว · ไม่สั่นผิดปกติ · หน้าแปลน/โบลต์/ที่ยึดไม่หลวม",
     "น้ำ (Drain) → สนิมในท่อ · กัดกร่อน · ท่อตัน\nเร่ง: ฝุ่น/สนิมจากงานเดินท่อหรือซ่อม",
     "อัตราไหลไม่พอ แรงดันตก · สนิมไปติดวาล์ว/กระบอกสูบ · Water hammer",
     "ท่อสั่น · หน้าแปลน/โบลต์/ที่ยึดหลวม: ทุกวัน\nลมรั่ว (Table 2): 3 เดือน",
     "—",
     "—",
     "5'-B-1, 5'-B-2, 5'-C-1, 5'-C-2 · EP26, EP28, EP29", "PDF p.41–42, p.44–45"),
    ("Air tube + Binder\n(สายลม + การรัด)",
     "ส่งลมระหว่างอุปกรณ์ · รัดเพื่อกันเสียดสีและให้ทำความสะอาดง่าย",
     "Air tube · Binder (เคเบิลไทร์) · Lock band (แบนด์ยึดสาย) · Band สำหรับรัดสายลม",
     "—",
     "สายกลม · รัด Binder เผื่อระยะ · Lock band ตรงขนาดและจำนวนสาย · ใช้ Band เฉพาะสายลม",
     "ไม่รัด → สายเสียดสีสึก\nเร่ง: รัดแน่นเกิน · Band ผิดขนาด/เกินจำนวน · ใช้ของทดแทน",
     "สายบุบ/แบน → แรงดันและอัตราไหลไม่พอ",
     "จุดรัด: ตา · รายสัปดาห์ *",
     "—",
     "—",
     "5'-C-4 · EP30", "PDF p.47"),
    ("Silencer / Exhaust port\n(ทางลมระบาย)",
     "ทางลมระบายของวาล์ว · จุดดูว่าน้ำมันหล่อลื่นไปถึง",
     "Silencer · Oil mist collector",
     "—",
     "มีน้ำมันที่ Silencer / Oil mist collector พอดี",
     "อุดตัน [INFER] · ละอองน้ำมันมากเกิน",
     "ละอองน้ำมันและ Drain ออกสู่บรรยากาศ ไม่ดีต่อสุขภาพ",
     "ดูน้ำมันที่ Silencer (แทนการนับหยด): ทุกวัน\nExhaust port: 3 เดือน",
     "—",
     "—",
     "5'-A-6, 5'-B-2, 5'-C-2 · EP25, EP26, EP29", "PDF p.40, p.42, p.45"),
    ("Compressor\n(แหล่งลม)",
     "ผลิตลมอัดให้ทั้งโรงงาน",
     "Compressor · Suction filter",
     "—",
     "[DATA REQUIRED] ตามมาตรฐานฝ่าย Utility",
     "น้ำมันหล่อลื่นของ Compressor ปนในลมเป็นละออง/ไอ (น้ำมันเสื่อม) · ฝุ่นเล็ดลอดช่อง Suction filter",
     "สิ่งปนเปื้อนเข้าระบบลมทั้งโรงงาน",
     "นอกขอบเขต AM ของเครื่อง → ดูแลโดยฝ่าย Utility [DATA REQUIRED]",
     "—",
     "—",
     "5'-A-2, 5'-B-2 · EP22, EP26", "PDF p.36, p.42"),
]

# 6 factors of failures (am-skill reference) → pneumatic examples
FORCED = [
    ("สภาพพื้นฐานไม่ครบ (Poor basic conditions: CLT)",
     "C: ไม่ระบาย Drain → น้ำเข้าระบบ · L: น้ำมันไม่หยด/หมด · T: Lock nut · หน้าแปลน · โบลต์ · ที่ยึดท่อหลวม",
     "PDF p.42–44, p.49", "ทำ CLT ตาม AM Check Sheet รายวัน/รายสัปดาห์", "Daily: Drain · หยดน้ำมัน · Lock nut · หน้าแปลน"),
    ("ไม่ใช้งานตามเงื่อนไข (Non-compliance with operating conditions)",
     "ค่าตั้งแรงดันเคลื่อน · หยดน้ำมันนอก 2–6 หยด/นาที · น้ำมันมากเกิน",
     "PDF p.39–40, p.45", "แถบเขียวที่เกจ · ป้ายค่าหยดน้ำมันที่ Lubricator *", "Daily: เกจแรงดัน · นับหยด"),
    ("ชิ้นส่วนเสื่อมไม่ได้แก้ (Deterioration not restored)",
     "Element ตัน · Diaphragm/สปริง/Packing/O-ring/Seal เสื่อม · Valve seat บาก",
     "PDF p.45–46", "ตรวจ 3 เดือน + ถอดตรวจ 1 ปี (PM *) · เปลี่ยนตามสภาพ", "Periodic: 3 เดือน / 1 ปี"),
    ("จุดอ่อนของการออกแบบ (Design weakness)",
     "ท่อยาวหรือเล็กเกิน → แรงดันตก · ตำแหน่งติดตั้ง Lubricator ไม่เหมาะ (น้ำมันไปไม่ถึง)",
     "PDF p.41, p.45", "แจ้ง PM / ทำ Kaizen sheet *", "Periodic: Exhaust port"),
    ("ใช้งาน/ซ่อมผิดวิธี (Mis-operation & mis-repair)",
     "เปิดถ้วย/ฝาน้ำมันขณะมีแรงดัน · รัด Binder แน่นเกิน · Lock band ผิดขนาด · ใช้ของทดแทน · ใช้ Element อัลลอยซ้ำ · ปรับหยดเอง",
     "PDF p.38, p.40, p.47", "สอน OPL (EP23, EP25, EP30) + Skill_Check", "Weekly: การรัดสายลม"),
    ("สภาพแวดล้อมเร่งการเสื่อม (Environment)",
     "กลางคืนอุณหภูมิลด → เกิดน้ำในระบบ · ฝุ่น/สนิมจากงานเดินท่อหรือซ่อม · น้ำมันเสื่อมจาก Compressor",
     "PDF p.42, p.44", "ระบาย Drain ต้นกะ * · ทำความสะอาดหลังงานท่อ * · แจ้ง Utility", "Daily: Drain"),
]

CONDITIONS = [
    ("ค่าตั้งแรงดัน Regulator", "[DATA REQUIRED] · ทำแถบเขียวที่ค่าตั้งบนเกจ", "PDF p.39"),
    ("ช่วงแรงดันใช้งานปกติ (เกจ)", "[DATA REQUIRED]", "PDF p.44"),
    ("หยดน้ำมัน Lubricator", "2–6 หยด/นาที (ตั้งตามอุปกรณ์และสภาพใช้งาน · ปรึกษาฝ่ายซ่อม)", "PDF p.40"),
    ("ชนิดน้ำมัน Lubricator", "[DATA REQUIRED] ตาม Catalog ผู้ผลิต", "PDF p.40"),
    ("เวลาต่อรอบกระบอกสูบ (ความเร็วปกติ)", "[DATA REQUIRED]", "PDF p.44"),
    ("รอบระบาย Drain", "วันละ 1 ครั้ง", "PDF p.44"),
    ("รอบเติมน้ำมัน Lubricator", "สัปดาห์ละ 1 ครั้ง · อัตราเดินเครื่องสูง ทุก 3 วัน", "PDF p.44"),
    ("รอบตรวจการทำงาน + ลมรั่ว", "ทุก 3 เดือน · ไม่พบปัญหา → ค่อยๆ ยืดรอบ", "PDF p.45"),
    ("รอบถอดตรวจภายใน", "ทุก 1 ปี", "PDF p.46"),
    ("Filter element", "โลหะ: ล้างน้ำมันก๊าด + แปรงไนลอน · อัลลอย (สีเงิน): เปลี่ยนใหม่ ห้ามใช้ซ้ำ", "PDF p.38"),
    ("ถ้วยกรอง", "ล้างด้วยน้ำหรือน้ำมันก๊าด แล้วเช็ดด้วยผ้าและแปรงไนลอน (ถ้วย Polycarbonate: ตรวจคู่มือผู้ผลิตก่อน *)", "PDF p.38"),
    ("สภาวะมาตรฐานของอากาศ (อ้างอิง)", "20 °C · 760 mmHg (ความดันสัมบูรณ์) · RH 65 %", "PDF p.35"),
]

QUIZ = [
    ("ชุด 3 ตัว (FRL) เรียงจากทางลมเข้าอย่างไร", "Air filter → Regulator → Lubricator", "EP31 · PDF p.48–49"),
    ("ต้องตรวจ Drain บ่อยแค่ไหน และทำไม", "วันละ 1 ครั้ง · กลางคืนอุณหภูมิลดจึงเกิดน้ำ", "EP28 · PDF p.44"),
    ("เข็มเกจ Regulator แกว่งมากขณะเดินเครื่อง หมายถึงอะไร ต้องทำอะไร", "ไส้กรอง Air filter ตัน → เปลี่ยน Filter element", "EP31 · PDF p.49"),
    ("Lock nut ของ Regulator หลวม จะเกิดอะไร", "Handle หมุนเองจากแรงสั่น → ค่าตั้งแรงดันเคลื่อน → ขันให้แน่น", "EP24, EP31 · PDF p.39, p.49"),
    ("มาตรฐานหยดน้ำมันของ Lubricator คือเท่าไร", "2–6 หยด/นาที · ปรึกษาฝ่ายซ่อมก่อนปรับ", "EP25 · PDF p.40"),
    ("น้ำมันในถ้วย Lubricator ขุ่นขาว หมายถึงอะไร ต้องทำอะไร", "มีน้ำปน → เปลี่ยนน้ำมันใหม่", "EP31 · PDF p.49"),
    ("น้ำมันใน Lubricator ลดเร็วผิดปกติ หมายถึงอะไร", "มีจุดรั่ว → หาจุดรั่วทันที", "EP28 · PDF p.44"),
    ("ก่อนถอดถ้วยกรอง หรือเปิดฝาเติมน้ำมัน ต้องทำอะไร เพราะอะไร", "ปล่อยแรงดันค้างให้หมด · ถ้วยอาจหลุดกระเด็น / น้ำมันพุ่งออก", "EP23, EP25 · PDF p.38, p.40"),
    ("รัดสายลมด้วย Binder แน่นเกินไป เกิดอะไร", "สายบุบ/แบน → แรงดันและอัตราไหลไม่พอ", "EP30 · PDF p.47"),
    ("หล่อลื่นไม่พอ ทำให้เกิดปัญหาอะไร (ตอบ 2 ข้อ)", "เสียดทานเพิ่ม (ไม่นิ่ง แรงไม่พอ วาล์วสลับไม่ดี) · ซีลสึก ลมรั่ว · สนิมติดในวาล์ว/ลูกสูบ", "EP27 · PDF p.43"),
]

CONFIRM = [
    ("ตำแหน่ง · Tag No. · จำนวนของแต่ละชิ้นส่วนบนเครื่องนี้", "ชีท System_Map", "[DATA REQUIRED]"),
    ("ค่าตั้งแรงดัน · ช่วงใช้งานปกติ · เวลาต่อรอบกระบอกสูบ · ชนิดน้ำมัน", "ชีท Operating_Conditions", "[DATA REQUIRED] · PDF p.39–40, p.44"),
    ("ผลกระทบที่ติด [INFER] (Regulator · เกจ · Speed controller · Silencer)", "ต้นฉบับไม่ระบุ ใช้หลักวิศวกรรมทั่วไป", "เกณฑ์ทั่วไป · ให้ฝ่ายซ่อมยืนยัน"),
    ("ผลกระทบของ Relief valve ในวงจรนี้", "ขึ้นกับตำแหน่งติดตั้ง", "[DATA REQUIRED]"),
    ("ผู้ทำงานถอดตรวจ 1 ปี (PM) และรอบรัดสายลมรายสัปดาห์", "ต้นฉบับไม่ระบุผู้ทำ/ความถี่", "ข้อเสนอ"),
    ("ตัดไฟ + ปล่อยลมค้างก่อนถอด Solenoid valve · LOTO", "ต้นฉบับเตือนเฉพาะถ้วยกรองและฝาน้ำมัน", "ข้อเสนอ · ให้ Safety ยืนยัน"),
    ("ระดับผ่าน Skill_Check ≥ 8/10", "ต้นฉบับไม่มีแบบทดสอบ", "ข้อเสนอ"),
    ("ขอบเขต Compressor (AM ของเครื่อง vs ฝ่าย Utility)", "ต้นฉบับอธิบายแต่ไม่แบ่งหน้าที่", "[DATA REQUIRED]"),
    ("คำแปลที่แก้: Bubble plate = Baffle plate · steam iron = กระบอกฉีดฝอย (霧吹き) · confusion = เพี้ยน (狂い) · sagging = หลวม (緩み)", "ใช้ในชีท Component_Analysis", "แก้คำแปล PDF p.38, p.40, p.46, p.49"),
]


# ---------------------------------------------------------------- helpers
def title(ws, text, sub, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, text).font = font(True, 15, "FFFFFF")
    ws.cell(1, 1).fill = fill(NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, sub).font = font(False, 9, "5A6775", italic=True)
    ws.cell(2, 1).alignment = Alignment(vertical="center", indent=1)


def header(ws, row, labels, widths):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row, i, lab)
        c.font, c.fill, c.alignment, c.border = font(True, 9, "FFFFFF"), fill(NAVY), CENTER, BOX
    ws.row_dimensions[row].height = 34
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, r, c, v, bold=False, align=WRAP, size=9):
    cell = ws.cell(r, c, v)
    tagged = isinstance(v, str) and ("[DATA REQUIRED]" in v or "[INFER]" in v or "*" in v)
    cell.font = font(bold, size, AMBER if tagged and not bold else INK)
    cell.alignment, cell.border = align, BOX
    if isinstance(v, str) and v.startswith("[DATA REQUIRED]"):
        cell.fill = fill(INPUT)
    return cell


def info(ws, row, pairs):
    for label, col in pairs:
        c = ws.cell(row, col, label)
        c.font, c.alignment = font(True, 9), Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + 2)
        for k in (col + 1, col + 2):
            ws.cell(row, k).fill, ws.cell(row, k).border = fill(INPUT), BOX
    ws.row_dimensions[row].height = 20


def page(ws, landscape=True, paper=8):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = paper
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3


wb = Workbook()

# ---------------------------------------------------------------- How_To_Use
ws = wb.active
ws.title = "How_To_Use"
title(ws, "Know My Machine — Pneumatic System (設備を知る = รู้จักเครื่องจักรของเรา · ระบบลม)",
      "AM Step 0 / Step 4 Module 4-4 (Pneumatics) · จาก OPL 5'-A-1 … 5'-C-5 (JIPM-Solutions · Equipment Skills Training – Pneumatic) · วิดีโอ EP22–EP31", 6)
info(ws, 3, [("เครื่องจักร", 1), ("Line / แผนก", 4)])
info(ws, 4, [("ผู้จัดทำ", 1), ("วันที่", 4)])
rows = [
    ("ชีท", "ใช้ทำอะไร"),
    ("System_Map", "ระบบลมของเครื่องตามทางลม 4 ส่วน (แหล่งลม · อุปกรณ์ประกอบ · ส่วนควบคุม · ส่วนทำงาน) + ช่องกรอกตำแหน่ง/Tag No./จำนวนบนเครื่องจริง"),
    ("Component_Analysis", "แต่ละชิ้นส่วน: หน้าที่ · โครงสร้าง · หลักการ · สภาพที่ถูกต้อง · การเสื่อม (ปกติ/เร่ง) · ผลถ้าปล่อยไว้ · จุดตรวจ AM · งาน PM · Safety"),
    ("Forced_Deterioration", "6 สาเหตุของการเสีย (6 Factors of failures) → ตัวอย่างในระบบลม → มาตรการ → ข้อตรวจใน AM Check Sheet"),
    ("Operating_Conditions", "ค่ามาตรฐานและรอบการดูแล (ตามต้นฉบับ) + ช่องกรอกค่าของเครื่องนี้"),
    ("Skill_Check", "แบบทดสอบ 10 ข้อหลังดูวิดีโอ + ตารางผลรายคน (ระดับทักษะ 1–4) · เฉลยอยู่ชีท Answer_Key"),
    ("To_Confirm", "ประเด็นที่หัวหน้างาน/ฝ่ายซ่อม/Safety ต้องยืนยันก่อนใช้จริง"),
    ("", ""),
    ("ป้ายที่มา", "PDF p.x = จากต้นฉบับ · [INFER] = หลักวิศวกรรมทั่วไป ไม่อยู่ในต้นฉบับ · * = ข้อเสนอ ให้หัวหน้างานยืนยัน · [DATA REQUIRED] = ค่าของเครื่อง (ช่องเหลือง)"),
    ("ความเชื่อมั่น (Confidence)", "HIGH: ข้อที่มี PDF p.x · MEDIUM: ข้อ [INFER] · ต้องกรอก [DATA REQUIRED] ก่อนใช้กับเครื่องจริง"),
    ("วิธีใช้", "1) ทีมเดินดูเครื่องจริงพร้อมชีท System_Map กรอกตำแหน่ง · 2) หัวหน้าทีมสอนด้วยวิดีโอ + Component_Analysis · 3) ทำ Skill_Check · 4) ติด Visual control ตามจุดตรวจ"),
    ("เชื่อมกับ", "AM_Check_Sheet_Pneumatic_System.xlsx (จุดตรวจรายวัน/สัปดาห์/รอบ) · วิดีโอ EP22–EP31"),
]
for i, (a, b) in enumerate(rows, start=6):
    head = i == 6
    ws.cell(i, 1, a).font = font(True, 10, "FFFFFF" if head else INK)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    ws.cell(i, 2, b).font = font(head, 10, "FFFFFF" if head else INK)
    for k in range(1, 7):
        ws.cell(i, k).border = BOX if a or b else Border()
        ws.cell(i, k).alignment = WRAPC
        if head:
            ws.cell(i, k).fill = fill(NAVY)
    ws.row_dimensions[i].height = 36 if a else 10
ws.column_dimensions["A"].width = 22
for col in "BCDEF":
    ws.column_dimensions[col].width = 24
page(ws, landscape=False, paper=9)

# ---------------------------------------------------------------- System_Map
ws = wb.create_sheet("System_Map")
cols = ["ส่วน (p.36)", "กลุ่ม", "ชิ้นส่วน", "หน้าที่", "ตำแหน่ง / Tag No. บนเครื่อง", "จำนวน", "OPL · วิดีโอ", "ที่มา"]
title(ws, "System Map — ระบบลมของเครื่อง (ตามทางลม)", "วงจรตัวอย่างตามต้นฉบับ PDF p.36 · ช่องสีเหลือง = กรอกจากการเดินดูเครื่องจริง", len(cols))
ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols))
c = ws.cell(3, 1, "ทางลม: " + FLOW)
c.font, c.fill, c.alignment = font(True, 10, "1F5FBF"), fill(SECTION), Alignment(vertical="center", indent=1, wrap_text=True)
ws.row_dimensions[3].height = 30
header(ws, 5, cols, [9, 26, 30, 44, 30, 9, 22, 16])
for r, row in enumerate(SYSTEM, start=6):
    order, group, comp, func, opl, src = row
    for k, v in enumerate((order, group, comp, func), start=1):
        put(ws, r, k, v, bold=(k == 3), align=CENTER if k == 1 else WRAPC)
    for k in (5, 6):
        put(ws, r, k, None, align=CENTER).fill = fill(INPUT)
    put(ws, r, 7, opl, align=WRAPC)
    put(ws, r, 8, src, align=WRAPC)
    ws.row_dimensions[r].height = 34
ws.freeze_panes = "C6"
ws.print_title_rows = "5:5"
page(ws)

# ---------------------------------------------------------------- Component_Analysis
ws = wb.create_sheet("Component_Analysis")
cols = ["ชิ้นส่วน", "หน้าที่ (Function)", "โครงสร้าง / ชิ้นส่วนสำคัญ", "หลักการทำงาน", "สภาพที่ถูกต้อง (Basic condition)",
        "การเสื่อม (ปกติ / เร่ง = Forced)", "ถ้าปล่อยไว้ → ผลกระทบ", "จุดตรวจ AM (วิธี · รอบ)", "งาน PM", "Safety", "OPL · วิดีโอ", "ที่มา"]
title(ws, "Know My Machine — Component Analysis · ระบบลม",
      "อ่านจากซ้ายไปขวา: รู้หน้าที่ → รู้สภาพที่ถูกต้อง → รู้ว่าเสื่อมอย่างไร → รู้ว่าต้องตรวจอะไร · ข้อความสีส้ม = [INFER] / * / [DATA REQUIRED]", len(cols))
header(ws, 4, cols, [17, 24, 26, 30, 26, 30, 26, 32, 24, 22, 17, 15])
for r, row in enumerate(COMPONENTS, start=5):
    for k, v in enumerate(row, start=1):
        cell = put(ws, r, k, v, bold=(k == 1))
        if k == 1:
            cell.fill = fill(SECTION)
        if k == 10 and v != "—":
            cell.font = font(True, 9, "C32A3E" if "[INFER]" not in v and "*" not in v else AMBER)
    ws.row_dimensions[r].height = 90
ws.freeze_panes = "B5"
ws.print_title_rows = "4:4"
page(ws)

# ---------------------------------------------------------------- Forced_Deterioration
ws = wb.create_sheet("Forced_Deterioration")
cols = ["6 สาเหตุของการเสีย (6 Factors of failures)", "ตัวอย่างในระบบลม", "ที่มา", "มาตรการ (Countermeasure)", "ข้อตรวจใน AM Check Sheet"]
title(ws, "Forced Deterioration — การเสื่อมแบบเร่ง (強制劣化 = เสื่อมเพราะคนหรือสภาพแวดล้อม)",
      "สมการ: การเสีย = ภาระ (Stress) > ความแข็งแรง (Strength) · 6 สาเหตุ → 6 มาตรการ (JIPM AM)", len(cols))
header(ws, 4, cols, [34, 56, 16, 40, 30])
for r, row in enumerate(FORCED, start=5):
    for k, v in enumerate(row, start=1):
        cell = put(ws, r, k, v, bold=(k == 1), align=WRAPC)
        if k == 1:
            cell.fill = fill(SECTION)
    ws.row_dimensions[r].height = 52
page(ws)

# ---------------------------------------------------------------- Operating_Conditions
ws = wb.create_sheet("Operating_Conditions")
cols = ["หัวข้อ", "ค่ามาตรฐาน / รอบ (ต้นฉบับ)", "ที่มา", "ค่าของเครื่องนี้", "ผู้ยืนยัน"]
title(ws, "Operating Conditions — ค่ามาตรฐานและรอบการดูแล", "ช่องสีเหลือง = กรอกค่าของเครื่องนี้ (ดู Catalog ผู้ผลิต / ฝ่ายซ่อม)", len(cols))
header(ws, 4, cols, [34, 60, 14, 28, 16])
for r, (a, b, src) in enumerate(CONDITIONS, start=5):
    put(ws, r, 1, a, bold=True, align=WRAPC)
    put(ws, r, 2, b, align=WRAPC)
    put(ws, r, 3, src, align=WRAPC)
    for k in (4, 5):
        put(ws, r, k, None, align=WRAPC).fill = fill(INPUT)
    ws.row_dimensions[r].height = 32
page(ws, landscape=False, paper=9)

# ---------------------------------------------------------------- Skill_Check
ws = wb.create_sheet("Skill_Check")
title(ws, "Skill Check — แบบทดสอบความรู้ระบบลม (10 ข้อ)", "ทำหลังดูวิดีโอ EP22–EP31 · เฉลยอยู่ชีท Answer_Key · ผ่าน ≥ 8/10 * (ข้อเสนอ)", 6)
info(ws, 3, [("ชื่อ", 1), ("วันที่", 4)])
header(ws, 5, ["ข้อ", "คำถาม", "คำตอบ", "", "", "ถูก (1/0)"], [6, 52, 30, 14, 14, 10])
for i, (q, _, ref) in enumerate(QUIZ, start=1):
    r = 5 + i
    put(ws, r, 1, i, align=CENTER)
    put(ws, r, 2, q + f"  ({ref.split(' · ')[0]})", align=WRAPC)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    for k in (3, 4, 5):
        ws.cell(r, k).border = BOX
    put(ws, r, 6, None, align=CENTER).fill = fill(INPUT)
    ws.row_dimensions[r].height = 40
score_row = 6 + len(QUIZ)
ws.merge_cells(start_row=score_row, start_column=1, end_row=score_row, end_column=5)
put(ws, score_row, 1, "คะแนนรวม", bold=True, align=Alignment(horizontal="right", vertical="center"))
for k in range(2, 6):
    ws.cell(score_row, k).border = BOX
sc = put(ws, score_row, 6, f"=SUM(F6:F{score_row - 1})", bold=True, align=CENTER)
sc.fill = fill(GREY)
dv01 = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
ws.add_data_validation(dv01)
dv01.add(f"F6:F{score_row - 1}")

# team result table
t0 = score_row + 2
ws.merge_cells(start_row=t0, start_column=1, end_row=t0, end_column=6)
ws.cell(t0, 1, "ผลรายคน (Skill matrix · Module 4-4 Pneumatics)").font = font(True, 11, "1F5FBF")
header(ws, t0 + 1, ["No.", "ชื่อ", "คะแนน (0–10)", "ผล", "ระดับทักษะ (1–4)", "ผู้ประเมิน"], [6, 52, 30, 14, 14, 10])
lv = DataValidation(type="list", formula1='"1 รู้,2 ทำได้,3 มั่นใจ,4 สอนได้"', allow_blank=True)
lv.prompt, lv.promptTitle = "1 รู้ (Know) · 2 ทำได้ (Can manage) · 3 มั่นใจ (Confident) · 4 สอนได้ (Can teach)", "ระดับทักษะ"
ws.add_data_validation(lv)
for n in range(1, 11):
    r = t0 + 1 + n
    put(ws, r, 1, n, align=CENTER)
    for k in (2, 3, 5, 6):
        put(ws, r, k, None, align=CENTER).fill = fill(INPUT)
    res = put(ws, r, 4, f'=IF(C{r}="","",IF(C{r}>=8,"ผ่าน","ทบทวน"))', bold=True, align=CENTER)
    res.fill = fill(GREY)
    lv.add(f"E{r}")
    ws.row_dimensions[r].height = 22
r = t0 + 12
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws.cell(r, 1, f'=COUNTIF(D{t0 + 2}:D{t0 + 11},"ผ่าน")&" คนผ่าน จาก "&COUNT(C{t0 + 2}:C{t0 + 11})&" คนที่ทดสอบ"').font = font(True, 10)
ws.cell(r, 1).alignment = Alignment(horizontal="right", vertical="center")
page(ws, landscape=False, paper=9)

# ---------------------------------------------------------------- Answer_Key
ws = wb.create_sheet("Answer_Key")
title(ws, "Answer Key — เฉลย Skill Check", "สำหรับผู้สอน · ไม่ต้องพิมพ์แจกผู้เรียน", 4)
header(ws, 4, ["ข้อ", "คำถาม", "เฉลย", "อ้างอิง"], [6, 50, 56, 22])
for i, (q, a, ref) in enumerate(QUIZ, start=1):
    r = 4 + i
    put(ws, r, 1, i, align=CENTER)
    put(ws, r, 2, q, align=WRAPC)
    put(ws, r, 3, a, bold=True, align=WRAPC)
    put(ws, r, 4, ref, align=WRAPC)
    ws.row_dimensions[r].height = 34
page(ws, landscape=False, paper=9)

# ---------------------------------------------------------------- To_Confirm
ws = wb.create_sheet("To_Confirm")
title(ws, "To Confirm — ประเด็นที่ต้องยืนยันก่อนใช้จริง", "ให้หัวหน้างาน / ฝ่ายซ่อมบำรุง / Safety ยืนยัน แล้วแก้ในชีทที่เกี่ยวข้อง", 7)
header(ws, 4, ["No.", "ประเด็น", "รายละเอียด", "ที่มา / สถานะ", "ผู้ยืนยัน", "วันที่", "ผล / ค่าที่ใช้"], [5, 50, 36, 30, 16, 12, 28])
for i, (a, b, c) in enumerate(CONFIRM, start=1):
    r = 4 + i
    for k, v in enumerate((i, a, b, c), start=1):
        put(ws, r, k, v, align=CENTER if k == 1 else WRAPC)
    for k in (5, 6, 7):
        put(ws, r, k, None).fill = fill(INPUT)
    ws.row_dimensions[r].height = 36
page(ws)

wb.save(OUT)
print("wrote", OUT)
